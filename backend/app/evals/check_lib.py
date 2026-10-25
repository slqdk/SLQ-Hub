#!/usr/bin/env python3
"""SLQ Hub tests: the checks that look at what a test case's turn made - run in the sandbox, never in the Hub, since
they open the model's files and run its code. Uses what the sandbox image has: Pillow, numpy, python-docx,
openpyxl, Tesseract, ffmpeg and the .NET SDK. Called in the case's working folder as

    python3 .hub/tests/check_lib.py <case id> .hub/tests/input.json

and prints one line: "@@CHECKS@@" followed by a JSON list of {name, ok, level, detail}. ok is null for a
measurement without a verdict (level "info"). The input names the turn's files and pictures, the case's own
input files and what a right result holds."""
from __future__ import annotations

import glob
import hashlib
import json
import os
import re
import subprocess
import sys

MARK = "@@CHECKS@@"
RESULTS: list[dict] = []
CASES: dict = {}


def check(name: str, ok, detail="", level: str = "must"):
    RESULTS.append({"name": name, "ok": None if ok is None else bool(ok), "level": level, "detail": str(detail)[:600]})
    return ok


def info(name: str, detail) -> None:
    check(name, None, detail, "info")


def case(name: str):
    def wrap(fn):
        CASES[name] = fn
        return fn
    return wrap


# ---------------------------------------------------------------- helpers

def run(cmd: list[str], timeout: int = 120, cwd: str | None = None, env: dict | None = None) -> tuple[int, str, str]:
    try:
        p = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout, cwd=cwd,
                           env={**os.environ, **(env or {})})
        return p.returncode, p.stdout, p.stderr
    except subprocess.TimeoutExpired as exc:
        out = exc.stdout or b""                    # bytes, even with text=True
        return 124, out.decode("utf-8", "replace") if isinstance(out, bytes) else out, f"timed out after {timeout} s"
    except OSError as exc:
        return 127, "", str(exc)


def mine(path: str) -> bool:
    """A file of the case's folder that the Hub did not put there for the checks or as an attachment."""
    return not path.startswith(".hub/")


def find(inp: dict, name: str | None = None, suffixes: tuple[str, ...] = ()) -> str | None:
    """A file the turn made: by name (anywhere in the folder) or by suffix; the turn's own list first, then a search
    of the folder (a file written by a script the turn ran is in the list too, but be generous)."""
    cands = [f["path"] for f in inp.get("files", [])]
    cands += [p[2:] if p.startswith("./") else p for p in glob.glob("./**/*", recursive=True)]
    for p in cands:
        if not mine(p) or not os.path.isfile(p):
            continue
        base = os.path.basename(p).lower()
        if name and base == name.lower():
            return p
    if suffixes:
        for p in cands:
            if mine(p) and os.path.isfile(p) and p.lower().endswith(suffixes):
                return p
    return None


def name_of(path: str | None) -> str:
    """A made file's name for a check's detail ("" for none, so the detail can say what is missing)."""
    return os.path.basename(path) if path else ""


def made_picture(inp: dict) -> str | None:
    """The picture the turn left for the user: its newest picture file in the folder (whichever way it was made -
    a tool, or its own code), else the picture tools' newest output."""
    pics = [f["path"] for f in inp.get("files", [])
            if str(f.get("mime", "")).startswith("image/") and os.path.isfile(f["path"])]
    return pics[0] if pics else output(inp, "images")


def output(inp: dict, kind: str) -> str | None:
    """The newest picture, sound or video the turn's tools made (their paths, from the tool results)."""
    items = [x for x in inp.get(kind, []) if os.path.isfile(x["path"])]
    return items[-1]["path"] if items else None


def probe(path: str) -> dict:
    code, out, err = run(["ffprobe", "-v", "error", "-print_format", "json", "-show_format", "-show_streams", path])
    if code != 0:
        raise RuntimeError(f"ffprobe: {err.strip()[:200]}")
    return json.loads(out)


def duration(meta: dict) -> float:
    try:
        return float(meta["format"]["duration"])
    except (KeyError, ValueError):
        return max((float(s.get("duration") or 0) for s in meta.get("streams", [])), default=0.0)


def loudness(path: str) -> tuple[float, float]:
    """Mean and peak volume in dB (ffmpeg volumedetect)."""
    code, out, err = run(["ffmpeg", "-nostdin", "-i", path, "-af", "volumedetect", "-f", "null", "-"], timeout=180)
    mean = re.search(r"mean_volume:\s*(-?[\d.]+|-inf) dB", err)
    peak = re.search(r"max_volume:\s*(-?[\d.]+|-inf) dB", err)
    f = (lambda m: float(m.group(1)) if m and m.group(1) != "-inf" else -200.0)
    return f(mean), f(peak)


def frame(path: str, at: float):
    """The video's picture at `at` seconds (Pillow)."""
    import io

    from PIL import Image
    p = subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-ss", f"{max(0.0, at):.3f}", "-i", path, "-frames:v",
                        "1", "-f", "image2pipe", "-vcodec", "png", "-"], capture_output=True, timeout=120)
    if p.returncode != 0 or not p.stdout:
        raise RuntimeError(f"no frame at {at:.1f} s: {p.stderr.decode('utf-8', 'replace')[:200]}")
    return Image.open(io.BytesIO(p.stdout)).convert("RGB")


def frame_jump(path: str) -> tuple[float, int]:
    """The biggest change from one frame to the next, at 64x36 grey, and where. A good clip changes by a point or
    two (the second run's kite: 1.3); the second run's broken clip - flat white, yellow and black every four frames
    - by 107."""
    import numpy as np
    p = subprocess.run(["ffmpeg", "-nostdin", "-v", "error", "-i", path, "-vf", "scale=64:36", "-f", "rawvideo",
                        "-pix_fmt", "gray", "-"], capture_output=True, timeout=300)
    fr = np.frombuffer(p.stdout, dtype=np.uint8)
    if p.returncode != 0 or fr.size < 2 * 64 * 36:
        return 0.0, 0
    fr = fr[: fr.size // (64 * 36) * 64 * 36].reshape(-1, 36, 64).astype(np.float32)
    d = np.abs(np.diff(fr, axis=0)).mean(axis=(1, 2))
    return float(d.max()), int(d.argmax()) + 1


def small_grey(img, size=(64, 36)):
    import numpy as np
    return np.asarray(img.convert("L").resize(size), dtype=np.float32)


def ocr_words(img) -> list[str]:
    """Words Tesseract reads in a picture, upper case - from the picture as it is and enlarged, light on dark too."""
    import tempfile

    from PIL import ImageOps
    grey = img.convert("L")
    variants = [img.convert("RGB"), grey.resize((grey.width * 2, grey.height * 2)), ImageOps.invert(grey)]
    words: list[str] = []
    with tempfile.TemporaryDirectory() as tmp:
        for i, v in enumerate(variants):
            path = os.path.join(tmp, f"v{i}.png")
            v.save(path)
            code, out, _ = run(["tesseract", path, "-", "-l", "eng+dan", "--psm", "11"], timeout=120)
            if code == 0:
                words += re.findall(r"[A-ZÆØÅ0-9]+", out.upper())
    return words


def distance(a: str, b: str) -> int:
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1]


def read_phrase(words: list[str], phrase: str) -> str | None:
    """Consecutive OCR words that spell `phrase` (spaces and dots aside), at most one letter in five off and its
    numbers exact: Tesseract reads a stylised headline in pieces and drops letters (on two posters spelled right it
    read "ULI HALLEN" and "JUL ALLEN")."""
    target = re.sub(r"[^A-ZÆØÅ0-9]", "", phrase.upper())
    best: tuple[int, str] | None = None
    for i in range(len(words)):
        joined = ""
        for j in range(i, min(i + 6, len(words))):
            joined += words[j]
            if len(joined) > len(target) + 3:
                break
            d = distance(joined, target)
            if re.sub(r"\D", "", joined) != re.sub(r"\D", "", target):
                continue                                # a number must be read exactly: 13. DECEMBER is wrong
            if best is None or d < best[0]:
                best = (d, " ".join(words[i:j + 1]))
    return best[1] if best is not None and best[0] <= max(1, len(target) // 5) else None


def unit_tests(test_file: str, module_dir: str, timeout: int = 180) -> tuple[int, int, str]:
    """Run a unittest file against modules in module_dir, in a process of its own: (tests run, problems, first
    problem)."""
    runner = (
        "import importlib.util, json, sys, unittest\n"
        "sys.path.insert(0, sys.argv[2])\n"
        "spec = importlib.util.spec_from_file_location('hidden_tests', sys.argv[1])\n"
        "mod = importlib.util.module_from_spec(spec)\n"
        "spec.loader.exec_module(mod)\n"
        "res = unittest.TextTestRunner(stream=open('/dev/null', 'w'), verbosity=0).run(\n"
        "    unittest.defaultTestLoader.loadTestsFromModule(mod))\n"
        "probs = res.failures + res.errors\n"
        "first = (probs[0][0].id().split('.')[-1] + ': ' + probs[0][1].strip().splitlines()[-1]) if probs else ''\n"
        "print('@@T@@' + json.dumps([res.testsRun, len(probs), first]))\n")
    code, out, err = run([sys.executable, "-c", runner, os.path.abspath(test_file), os.path.abspath(module_dir)],
                         timeout=timeout)
    m = re.search(r"@@T@@(.*)", out)
    if not m:
        tail = (err.strip().splitlines() or ["no output"])[-1]
        return 0, 1, f"the tests did not run: {tail[:300]}"
    n, bad, first = json.loads(m.group(1))
    return n, bad, first


def red_blue(img):
    """Shares of strongly red and strongly blue pixels."""
    import numpy as np
    a = np.asarray(img.convert("RGB"), dtype=np.float32)
    r, g, b = a[..., 0], a[..., 1], a[..., 2]
    red = (r > 110) & (r > 1.8 * g) & (r > 1.8 * b)
    blue = (b > 100) & (b > 1.35 * r) & (b > 1.15 * g)
    return red, blue


# ---------------------------------------------------------------- documents

@case("documents.csv")
def _sales(inp: dict) -> None:
    p = find(inp, "omsaetning.png") or find(inp, None, (".png",))
    if not check("the chart was saved as omsaetning.png", p is not None and os.path.basename(p) == "omsaetning.png",
                 p or "no PNG made"):
        if p is None:
            return
    from PIL import Image
    with Image.open(p) as img:
        check("the chart is a picture of a sensible size", img.width >= 400 and img.height >= 250,
              f"{p}: {img.width}x{img.height} px", "should")


@case("documents.word")
def _word(inp: dict) -> None:
    p = find(inp, "rapport.docx") or find(inp, None, (".docx",))
    if not check("rapport.docx was made", p is not None, p or "no .docx in the folder"):
        return
    import docx
    d = docx.Document(p)
    texts = [x.text.strip() for x in d.paragraphs]
    check("the title is there", any("månedsrapport oktober 2026" in t.lower() for t in texts) or
          "månedsrapport oktober 2026" in (d.core_properties.title or "").lower(), texts[:3])
    heads = [(i, x.text.strip()) for i, x in enumerate(d.paragraphs)
             if (x.style is not None and x.style.name.lower().startswith(("heading", "overskrift")))]
    wanted = ("Produktion", "Kvalitet", "Plan for november")
    found = {w: next((i for i, t in heads if t.lower() == w.lower()), None) for w in wanted}
    check("the three headings are real headings", all(v is not None for v in found.values()),
          f"headings: {[t for _, t in heads][:8]}")
    tables = d.tables
    shapes = [(len(t.rows), len(t.columns)) for t in tables]
    check("a table with 4 rows and 3 columns", (4, 3) in shapes, f"tables: {shapes}")
    # the body in order (its paragraphs are d.paragraphs, in the same order): where the table is
    order, idx = [], 0
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            order.append(("p", idx))
            idx += 1
        elif tag == "tbl":
            order.append(("t", None))
    prod, qual, plan = (found.get(w) for w in wanted)
    if prod is not None and qual is not None:
        seen_p, ok = None, False
        for kind, k in order:
            if kind == "p":
                seen_p = k
            elif seen_p is not None and prod <= seen_p < qual:
                ok = True
        check("the table is under Produktion", ok, "", "should")
    if plan is not None:
        items = 0
        for x in d.paragraphs[plan + 1:]:
            if x.style is not None and x.style.name.lower().startswith(("heading", "overskrift")):
                break
            numbered = x._p.pPr is not None and x._p.pPr.numPr is not None
            if x.text.strip() and (numbered or "list" in (x.style.name.lower() if x.style is not None else "")):
                items += 1
        check("a list of 3 points under Plan for november", items >= 3, f"{items} list items", "should")


@case("documents.excel")
def _excel(inp: dict) -> None:
    p = find(inp, "budget.xlsx") or find(inp, None, (".xlsx",))
    if not check("budget.xlsx was made", p is not None, p or "no .xlsx in the folder"):
        return
    import openpyxl
    wb = openpyxl.load_workbook(p)                 # formulas as written, not their cached values
    ws = wb.active
    for sheet in wb.worksheets:
        if str(sheet["A1"].value or "").strip().lower() == "måned":
            ws = sheet
            break
    heads = [str(ws.cell(1, c).value or "").strip().lower() for c in range(1, 5)]
    check("the headings in row 1", heads == ["måned", "indtægter", "udgifter", "resultat"], heads)
    months = [str(ws.cell(r, 1).value or "").strip().lower() for r in range(2, 14)]
    names = ["januar", "februar", "marts", "april", "maj", "juni", "juli", "august", "september", "oktober",
             "november", "december"]
    check("January to December in A2:A13", months == names or all(m[:3] == n[:3] for m, n in zip(months, names)),
          months, "should")
    bad = []
    for r in range(2, 14):
        f = ws.cell(r, 4).value
        if not (isinstance(f, str) and f.startswith("=") and re.search(rf"B\$?{r}\b", f.upper())
                and re.search(rf"C\$?{r}\b", f.upper())):
            bad.append(f"D{r}={f!r}")
    check("Resultat is a formula of B and C in every month", not bad, ", ".join(bad[:4]) or "D2:D13 are formulas")
    totals = [ws.cell(14, c).value for c in (2, 3, 4)]
    check("row 14 sums B, C and D with SUM", all(isinstance(t, str) and "SUM(" in t.upper() for t in totals),
          f"B14:D14 = {totals}")
    nums = [ws.cell(r, c).value for r in range(2, 14) for c in (2, 3)]
    check("numbers in Indtægter and Udgifter", all(isinstance(v, (int, float)) for v in nums),
          f"{sum(isinstance(v, (int, float)) for v in nums)} of 24 cells are numbers", "should")


# ---------------------------------------------------------------- pictures

@case("images.poster")
def _poster(inp: dict) -> None:
    p = output(inp, "images")
    if not check("a picture was made", p is not None, name_of(p) or "no picture from the picture tools"):
        return
    from PIL import Image
    img = Image.open(p).convert("RGB")
    info("size", f"{img.width}x{img.height} px")
    words = ocr_words(img)
    for target in ("JUL I HALLEN", "12. DECEMBER"):
        got = read_phrase(words, target)
        check(f"Tesseract reads {target}", got is not None,
              f"read as {got}" if got else f"read: {' '.join(words[:30])}"[:300], "should")
    import numpy as np
    a = np.asarray(img, dtype=np.float32)
    h, w = a.shape[:2]
    m = max(4, int(min(h, w) * 0.06))
    border = np.concatenate([a[:m].reshape(-1, 3), a[-m:].reshape(-1, 3), a[:, :m].reshape(-1, 3),
                             a[:, -m:].reshape(-1, 3)])
    r, g, b = border.mean(0)
    check("a dark blue background", b > r + 12 and b >= g and (0.3 * r + 0.59 * g + 0.11 * b) < 120,
          f"mean colour at the edges: rgb({r:.0f}, {g:.0f}, {b:.0f})", "should")


@case("images.edit")
def _edit(inp: dict) -> None:
    p = made_picture(inp)
    src = inp.get("inputs", {}).get("picture")
    if not check("an edited picture was made", p is not None, name_of(p) or "no new picture"):
        return
    import numpy as np
    from PIL import Image
    a = Image.open(src).convert("RGB")
    b = Image.open(p).convert("RGB")
    info("size", f"{b.width}x{b.height} px (the original: {a.width}x{a.height})")
    red0, _ = red_blue(a)
    red1, blue1 = red_blue(b)
    share0, r1, b1 = red0.mean(), red1.mean(), blue1.mean()
    check("the ball is blue now", b1 >= 0.5 * share0, f"blue {b1:.1%} of the picture (the red ball was {share0:.1%})")
    check("no red ball left", r1 <= 0.3 * share0, f"red {r1:.1%}", "should")
    exp = inp.get("expected", {})
    bb = b.resize(a.size)
    ax, bx = np.asarray(a, dtype=np.float32), np.asarray(bb, dtype=np.float32)
    yy, xx = np.mgrid[0:a.height, 0:a.width]
    cx, cy = exp.get("center", (a.width / 2, a.height / 2))
    rad = exp.get("radius", 80)
    away = ((xx - cx) ** 2 + (yy - cy) ** 2) > (rad * 1.6) ** 2
    diff = np.abs(ax - bx).mean(-1)[away].mean()
    check("the rest of the picture is unchanged", diff < 28, f"mean difference away from the ball: {diff:.1f} of 255",
          "should")


@case("images.cutout")
def _cutout(inp: dict) -> None:
    p = made_picture(inp)
    if not check("a cut-out was made", p is not None, name_of(p) or "no new picture"):
        return
    import numpy as np
    from PIL import Image
    img = Image.open(p)
    if not check("it has a transparent background (alpha)", img.mode in ("RGBA", "LA", "PA") or
                 "transparency" in img.info, f"mode {img.mode}"):
        return
    a = np.asarray(img.convert("RGBA"))[..., 3].astype(np.float32)
    h, w = a.shape
    clear = (a < 32).mean()
    check("most of it is transparent, not all", 0.5 <= clear <= 0.985, f"{clear:.1%} transparent")
    exp = inp.get("expected", {})
    sw, sh = exp.get("size", (w, h))
    cx, cy = exp.get("center", (sw / 2, sh / 2))
    x, y = int(cx * w / sw), int(cy * h / sh)
    check("the ball is kept", a[y, x] > 200, f"alpha at the ball's centre: {a[y, x]:.0f}", "should")
    corners = [a[2, 2], a[2, w - 3], a[h - 3, 2], a[h - 3, w - 3]]
    check("the corners are transparent", max(corners) < 32, f"alpha in the corners: {[int(c) for c in corners]}",
          "should")


# ---------------------------------------------------------------- sound

def _sound(inp: dict, what: str) -> None:
    p = output(inp, "audio")
    if not check(f"{what} was made", p is not None, name_of(p) or "no sound from the sound tools"):
        return
    meta = probe(p)
    secs = duration(meta)
    exp = inp.get("expected", {})
    lo, hi = exp.get("seconds", (0, 10_000))
    check(f"it lasts {lo}-{hi} s", lo <= secs <= hi, f"{secs:.1f} s ({os.path.basename(p)})", exp.get("level", "must"))
    mean, peak = loudness(p)
    check("it is not silent", peak > -30, f"mean {mean:.1f} dB, peak {peak:.1f} dB")
    if peak >= -0.1:
        check("it does not clip", False, f"peak {peak:.1f} dB", "should")


@case("sound.speech")
def _speech(inp: dict) -> None:
    _sound(inp, "speech")


@case("sound.music")
def _music(inp: dict) -> None:
    _sound(inp, "music")


@case("sound.effect")
def _effect(inp: dict) -> None:
    _sound(inp, "a sound effect")


# ---------------------------------------------------------------- video

def _video(inp: dict) -> dict | None:
    p = output(inp, "video")
    if not check("a video was made", p is not None, name_of(p) or "no video from video_generate"):
        return None
    meta = probe(p)
    v = next((s for s in meta.get("streams", []) if s.get("codec_type") == "video"), {})
    secs = duration(meta)
    exp = inp.get("expected", {})
    lo, hi = exp.get("seconds", (2, 11))
    check(f"it lasts {lo}-{hi} s", lo <= secs <= hi, f"{secs:.2f} s")
    w, h = int(v.get("width") or 0), int(v.get("height") or 0)
    rate = v.get("avg_frame_rate") or "0/1"
    num, _, den = rate.partition("/")
    fps = float(num) / float(den or 1) if float(den or 1) else 0.0
    check("at least 640 pixels wide", w >= 640, f"{w}x{h}, {fps:.0f} fps, {v.get('codec_name')}", "should")
    first, last = frame(p, 0.0), frame(p, max(0.0, secs - 0.25))
    moved = float(abs(small_grey(first) - small_grey(last)).mean())
    check("it moves", moved > 2.0, f"difference first-last frame: {moved:.1f} of 255", "should")
    jump, at = frame_jump(p)
    check("it moves smoothly (no jumps between frames)", jump <= 30,
          f"biggest change from one frame to the next: {jump:.1f} of 255 (frame {at})", "should")
    return {"path": p, "first": first, "seconds": secs}


@case("video.t2v")
def _t2v(inp: dict) -> None:
    _video(inp)


@case("video.i2v")
def _i2v(inp: dict) -> None:
    got = _video(inp)
    src = inp.get("inputs", {}).get("picture")
    if not got or not src:
        return
    from PIL import Image
    diff = float(abs(small_grey(Image.open(src)) - small_grey(got["first"])).mean())
    check("it starts from the picture", diff < 35, f"difference picture-first frame: {diff:.1f} of 255", "should")


# ---------------------------------------------------------------- code

@case("code.python")
def _roman(inp: dict) -> None:
    p = find(inp, "roman.py")
    if not check("roman.py was written", p is not None, p or "no roman.py in the folder"):
        return
    n, bad, first = unit_tests(".hub/tests/test_roman_hidden.py", os.path.dirname(p) or ".")
    check("hidden tests pass", n > 0 and bad == 0, f"{n - bad} of {n} passed" + (f"; {first}" if first else ""))
    own = [f["path"] for f in inp.get("files", []) if re.search(r"(^|/)test[^/]*\.py$", f["path"])]
    check("it wrote tests of its own", bool(own), ", ".join(own) or "no test file", "should")


@case("code.debug")
def _inventory(inp: dict) -> None:
    p = "lager/inventory.py" if os.path.isfile("lager/inventory.py") else find(inp, "inventory.py")
    if not check("inventory.py is there", p is not None, p or "lager/inventory.py is gone"):
        return
    folder = os.path.dirname(p) or "."
    n, bad, first = unit_tests(".hub/tests/test_inventory_orig.py", folder)
    check("the given tests pass", n > 0 and bad == 0, f"{n - bad} of {n} passed" + (f"; {first}" if first else ""))
    n, bad, first = unit_tests(".hub/tests/test_inventory_hidden.py", folder)
    check("hidden tests pass", n > 0 and bad == 0, f"{n - bad} of {n} passed" + (f"; {first}" if first else ""),
          "should")
    t = os.path.join(folder, "test_inventory.py")
    same = os.path.isfile(t) and hashlib.sha256(open(t, "rb").read()).hexdigest() == inp["expected"]["tests_sha"]
    check("the tests were not changed", same, "" if same else "test_inventory.py differs from the one given", "should")


def _top(out: str) -> list[list]:
    return [[m.group(1).lower(), int(m.group(2))] for m in
            (re.match(r"^\s*([^\W\d_]+)\s*:\s*(\d+)\s*$", line) for line in out.splitlines()) if m][:5]


@case("code.csharp")
def _words(inp: dict) -> None:
    proj = sorted(glob.glob("ordtaeller/**/*.csproj", recursive=True)) or \
        sorted(p for p in glob.glob("**/*.csproj", recursive=True) if mine(p))
    if not check("a C# project in ordtaeller", bool(proj) and proj[0].startswith("ordtaeller/"),
                 proj[0] if proj else "no .csproj"):
        if not proj:
            return
    folder = os.path.dirname(proj[0])
    exp = inp["expected"]
    for label, text, level in (("tekst.txt", "tekst.txt", "must"), ("a text it has not seen", ".hub/tests/skjult.txt",
                                                                     "should")):
        code, out, err = run(["dotnet", "run", "--project", folder, "--", os.path.abspath(text)], timeout=420)
        if code != 0:
            check(f"it runs on {label}", False, (err or out).strip()[-400:], level)
            continue
        want = exp["visible" if text == "tekst.txt" else "hidden"]
        got = _top(out)
        check(f"the top 5 for {label}", got == want, f"got {got}, expected {want}", level)


# ---------------------------------------------------------------- tasks

@case("agent.files")
def _files(inp: dict) -> None:
    data = "maalinger/data.csv" if os.path.isfile("maalinger/data.csv") else find(inp, "data.csv")
    script = find(inp, "analyse.py")
    result = "maalinger/resultat.md" if os.path.isfile("maalinger/resultat.md") else find(inp, "resultat.md")
    check("maalinger/data.csv", data == "maalinger/data.csv", data or "none")
    check("analyse.py", script is not None, script or "none")
    if not check("resultat.md", result is not None, result or "none") or data is None:
        return
    import csv
    with open(data, encoding="utf-8-sig") as f:
        rows = [r for r in csv.reader(f, delimiter=";" if ";" in f.readline() else ",")]
    temps = []
    for r in rows:
        try:
            temps.append(float(r[1].replace(",", ".")))
        except (IndexError, ValueError):
            pass
    check("24 measurements between 15 and 25 degrees", len(temps) == 24 and all(15 <= t <= 25 for t in temps),
          f"{len(temps)} values, {min(temps, default=0)}-{max(temps, default=0)}")
    if not temps:
        return
    mean, lo, hi = sum(temps) / len(temps), min(temps), max(temps)
    text = open(result, encoding="utf-8", errors="replace").read()
    nums = [float(x.replace(",", ".")) for x in re.findall(r"-?\d+(?:[.,]\d+)?", text)]
    missing = [f"{name} {v:.2f}" for name, v in (("mean", mean), ("min", lo), ("max", hi))
               if not any(abs(n - v) < 0.006 for n in nums)]
    check("resultat.md has the right mean, min and max", not missing,
          f"missing: {', '.join(missing)}" if missing else f"mean {mean:.2f}, min {lo:.2f}, max {hi:.2f}")
    said = [float(x.replace(",", ".")) for x in re.findall(r"-?\d+(?:[.,]\d+)?", inp.get("answer", ""))]
    check("the answer gives the mean", any(abs(n - mean) < 0.006 for n in said), f"mean {mean:.2f}", "should")


def main(argv: list[str]) -> int:
    name, path = argv[1], argv[2]
    with open(path, encoding="utf-8") as f:
        inp = json.load(f)
    fn = CASES.get(name)
    if fn is None:
        check("the check exists", False, f"no sandbox check for {name}")
    else:
        try:
            fn(inp)
        except Exception as exc:  # noqa: BLE001 - reported as a failed check, with what went wrong
            check("the checks ran", False, f"{type(exc).__name__}: {exc}")
    print(MARK + json.dumps(RESULTS, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
