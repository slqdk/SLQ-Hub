"""Working with the files of the project (or the chat's scratch folder)."""
from __future__ import annotations

import io
import re
from difflib import SequenceMatcher
from fractions import Fraction
from pathlib import Path

from ...chats import store as chats
from ...projects.store import is_text, mime_of
from .base import Tool, ToolContext, ToolResult, clip, obj, register

MAX_READ_LINES = 2000
MAX_READ_CHARS = 120_000
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


def _fmt_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n} B"


async def list_files(ctx: ToolContext, args: dict) -> ToolResult:
    wd = ctx.workdir
    base = wd.resolve(args.get("path") or ".")
    if not base.exists():
        return ToolResult(f"{args.get('path')} does not exist.", is_error=True)
    depth = max(1, min(int(args.get("depth") or 3), 8))
    pattern = args.get("pattern")
    lines, count = [], 0
    root_depth = len(base.relative_to(wd.root.resolve()).parts) if base != wd.root.resolve() else 0
    for e in wd.tree(wd.rel(base)):
        p = Path(e["path"])
        if len(p.parts) - root_depth > depth:
            continue
        if pattern and not e["dir"] and not p.match(pattern):
            continue
        indent = "  " * (len(p.parts) - root_depth - 1)
        lines.append(f"{indent}{p.name}/" if e["dir"] else f"{indent}{p.name}  ({_fmt_size(e['size'])})")
        count += 1
        if count >= 500:
            lines.append("... (more entries not shown - narrow the path or pattern)")
            break
    head = f"Folder {wd.sandbox_path}/{wd.rel(base) if wd.rel(base) != '.' else ''}".rstrip("/")
    return ToolResult(head + ":\n" + ("\n".join(lines) if lines else "(empty)"))


def _pdf_text(data: bytes, first: int, last: int | None) -> tuple[str, int]:
    import pymupdf
    doc = pymupdf.open(stream=data, filetype="pdf")
    n = doc.page_count
    last = min(n, last or n)
    parts = [f"--- page {i + 1} ---\n{doc[i].get_text()}" for i in range(max(0, first - 1), last)]
    return "\n".join(parts), n


def _docx_text(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)


MAX_SHEET_ROWS = 2000


def _xlsx_text(data: bytes) -> str:
    """Every sheet as rows of cells separated by |, each cell's value and - where it is one - its formula (0.23.0:
    skills are made from Excel files too, also by people whose chats run no code)."""
    import openpyxl
    values = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    out = []
    for ws_v, ws_f in zip(values.worksheets, formulas.worksheets):
        out.append(f"--- sheet {ws_v.title} ---")
        shown = 0
        for row_v, row_f in zip(ws_v.iter_rows(), ws_f.iter_rows()):
            cells = []
            for cv, cf in zip(row_v, row_f):
                v, f = cv.value, cf.value
                text = "" if v is None else str(v)
                if isinstance(f, str) and f.startswith("="):
                    text = f"{text} ({f})" if text else f
                cells.append(text)
            while cells and not cells[-1]:
                cells.pop()
            if not cells:
                continue
            if shown >= MAX_SHEET_ROWS:
                out.append(f"(more rows - the sheet goes on past row {row_v[0].row - 1})")
                break
            out.append(f"{row_v[0].row}: " + " | ".join(cells))
            shown += 1
    return "\n".join(out)


def _pptx_text(data: bytes) -> str:
    """Every slide's text, its tables and the speaker notes."""
    from pptx import Presentation
    deck = Presentation(io.BytesIO(data))
    out = []
    for i, slide in enumerate(deck.slides, start=1):
        out.append(f"--- slide {i} ---")
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"Notes: {notes}")
    return "\n".join(out)


DOCUMENTS = {".docx": ("Word document", _docx_text), ".xlsx": ("Excel workbook", _xlsx_text),
             ".xlsm": ("Excel workbook", _xlsx_text), ".pptx": ("PowerPoint presentation", _pptx_text)}


def document_text(name: str, data: bytes) -> tuple[str, str] | None:
    """(what it is, its text) for a Word, Excel or PowerPoint file; None for other files. A file that cannot be
    read says so instead of raising."""
    kind = DOCUMENTS.get(Path(name).suffix.lower())
    if kind is None:
        return None
    try:
        return kind[0], kind[1](data)
    except Exception as exc:  # noqa: BLE001 - a damaged or protected file
        return kind[0], f"(could not be read: {exc.__class__.__name__}: {str(exc)[:200]})"


async def read_file(ctx: ToolContext, args: dict) -> ToolResult:
    rel = args.get("path") or ""
    p = ctx.workdir.resolve(rel)
    if not p.is_file():
        return ToolResult(f"No such file: {rel}", is_error=True)
    data = p.read_bytes()
    ext = p.suffix.lower()
    if ext in IMAGE_EXT:
        up = chats.save_upload(data, p.name, mime_of(p.name))
        return ToolResult(f"The image {rel} ({up.get('width')}x{up.get('height')} px) is shown below.",
                          images=[{"sha": up["sha"], "mime": up["mime"]}])
    if ext == ".pdf":
        first = int(args.get("offset") or 1)
        text, n = _pdf_text(data, first, int(args["limit"]) + first - 1 if args.get("limit") else None)
        text, cut = clip(text, MAX_READ_CHARS)
        return ToolResult(f"{rel}: PDF with {n} pages (offset/limit select pages).\n{text}")
    doc = document_text(p.name, data)
    if doc is not None:
        text, cut = clip(doc[1], MAX_READ_CHARS)
        return ToolResult(f"{rel} ({doc[0]} text):\n{text}")
    if not is_text(data):
        return ToolResult(f"{rel} is a binary file ({_fmt_size(len(data))}, {mime_of(p.name)}). "
                          "Use bash with a suitable tool to inspect it.", is_error=True)
    lines = data.decode("utf-8", "replace").splitlines()
    offset = max(1, int(args.get("offset") or 1))
    limit = max(1, min(int(args.get("limit") or MAX_READ_LINES), MAX_READ_LINES))
    chunk = lines[offset - 1: offset - 1 + limit]
    width = len(str(offset + len(chunk)))
    body = "\n".join(f"{str(i).rjust(width)}\t{l}" for i, l in enumerate(chunk, start=offset))
    body, _ = clip(body, MAX_READ_CHARS)
    more = ""
    if offset - 1 + limit < len(lines):
        more = f"\n(lines {offset}-{offset + len(chunk) - 1} of {len(lines)}; read on with offset={offset + len(chunk)})"
    return ToolResult(f"{rel}:\n{body}{more}" if body else f"{rel} is empty.")


def copied_mark(new: str, there=lambda: "") -> str | None:
    """A mark the Hub puts where it left text out of what the model was shown, in text it writes and not already in
    the file (there: its text, or what gives it - read only when there is a mark): a copy of shortened context, not
    content (chat 79: its game's index.html got a heading cut short and "... [shortened by 24 characters]"). The
    mark, or None."""
    from ..context import MARK
    found = [m.group(0) for m in MARK.finditer(new)]
    if not found:
        return None
    text = there() if callable(there) else there
    return next((f for f in found if f not in text), None)


def marked(mark: str, rel: str) -> ToolResult:
    return ToolResult(f"Not written: the text for {rel} has \"{mark}\" in it. That is the Hub's mark for text left out "
                      "of what you were shown - it belongs to no file. Write that part in full (read the file or the "
                      "source again if you need its text).", is_error=True)


async def write_file(ctx: ToolContext, args: dict) -> ToolResult:
    rel = args.get("path") or ""
    content = args.get("content")
    if not isinstance(content, str):
        return ToolResult("content must be a string", is_error=True)
    p = ctx.workdir.resolve(rel)
    existed = p.exists()
    mark = copied_mark(content, lambda: p.read_text(encoding="utf-8", errors="replace") if p.is_file() else "")
    if mark:
        return marked(mark, rel)
    # Files this turn writes in parts. Chat 21 (0.20.2): twice the model left out append: true for the next part,
    # replaced the whole game with it and spent 12 of its 60 steps building the file again.
    in_parts = ctx.state.setdefault("written_in_parts", set())
    key = ctx.workdir.rel(p)
    if args.get("append") is True:
        in_parts.add(key)
    if args.get("append") is True and existed:       # the next part of a long file (0.16.2)
        if not p.is_file():
            return ToolResult(f"{rel} is a folder, not a file", is_error=True)
        data = p.read_bytes() + content.encode()
        ctx.workdir.write(rel, data, "agent", "hub", f"agent added to {rel}")
        lines = data.count(b"\n") + 1
        return ToolResult(f"Added {len(content.encode())} bytes to {rel} (now {len(data)} bytes, {lines} lines).")
    if existed and key in in_parts and args.get("replace") is not True and p.is_file():
        return ToolResult(f"Not written: you are writing {rel} in parts in this turn, and this call has no append: "
                          f"true - it would replace the {p.stat().st_size:,} bytes there with these "
                          f"{len(content.encode()):,}. If this is the next part, call write_file again with append: "
                          "true. To replace the whole file, call it with replace: true.", is_error=True)
    ctx.workdir.write(rel, content, "agent", "hub", f"agent wrote {rel}")
    return ToolResult(f"{'Overwrote' if existed else 'Created'} {rel} ({len(content.encode())} bytes, "
                      f"{content.count(chr(10)) + 1} lines).")


NUMBERED = re.compile(r"^ *\d+\t")          # the line number read_file puts before each line ("  12\t")


def _unnumbered(s: str, every: bool = True) -> str | None:
    """s without read_file's line numbers, when every line with text in it has one: a copy of what read_file showed,
    numbers and all. None when s is not numbered. every=False: the numbers left out where there are any."""
    lines = s.split("\n")
    if every and (not any(l.strip() for l in lines) or not all(NUMBERED.match(l) for l in lines if l.strip())):
        return None
    return "\n".join(NUMBERED.sub("", l, count=1) for l in lines)


def _squeeze(line: str) -> str:
    return " ".join(line.split())


def _indent(line: str) -> str:
    return line[: len(line) - len(line.lstrip(" \t"))]


def _step(table: dict[str, str]):
    """How one step of old's indentation is written in the file, as the table of old's indentations and the file's
    shows it (four spaces for a tab, two spaces for four): a function for the part of an indentation past one the
    table has. It leaves that part as it is when the table does not show a single, regular step."""
    steps = {(a2[len(a1):], b2[len(b1):]) for a1, b1 in table.items() for a2, b2 in table.items()
             if len(a2) > len(a1) and len(b2) > len(b1) and a2.startswith(a1) and b2.startswith(b1)}
    steps = steps or {(a, b) for a, b in table.items() if a and b}
    kinds = {(a[0], b[0], Fraction(len(b), len(a))) for a, b in steps}
    if not steps or len(kinds) != 1 or any(len(set(a)) > 1 or len(set(b)) > 1 for a, b in steps):
        return lambda rest: rest
    (was, now, ratio), = kinds

    def convert(rest: str) -> str:
        if rest and set(rest) == {was} and (len(rest) * ratio).denominator == 1:
            return now * int(len(rest) * ratio)
        return rest
    return convert


def _fit_indent(new: list[str], old: list[str], there: list[str]) -> list[str] | None:
    """new indented the way the file is, from what old's lines show of the difference: each indentation in old stands
    for the file's on the same line (four spaces for a tab, a first line copied without its indentation). None when
    that is not one-to-one, or new has an indentation it cannot place."""
    pairs = [(_indent(a), _indent(b)) for a, b in zip(old, there) if a.strip()]
    if all(a == b for a, b in pairs):
        return new
    table: dict[str, str] = {}
    for a, b in pairs[1:]:
        if table.setdefault(a, b) != b:
            return None
    first_a, first_b = pairs[0]
    alone = table.get(first_a, first_b) != first_b       # the first line alone is off: copied from after its indent
    if not alone:
        table[first_a] = first_b
    step, out = _step(table), []
    for j, line in enumerate(new):
        if not line.strip():
            out.append("")
            continue
        ind = _indent(line)
        if j == 0 and alone and ind == first_a:
            out.append(first_b + line[len(ind):])
            continue
        keys = [k for k in table if ind.startswith(k)]
        if not keys:
            return None
        k = max(keys, key=len)                          # a deeper line: the known part, then the rest converted
        out.append(table[k] + step(ind[len(k):]) + line[len(ind):])
    return out


def _loose_edit(text: str, old: str, new: str) -> tuple[str, int, bool] | str | None:
    """old in the file when whitespace does not count - line by line, the spaces and tabs at the ends of a line left
    out and a run of them inside counted as one - for a copy that came out with tabs as spaces, an indentation lost or
    a trailing space. In exactly one place: (the text with new there, its first line, whether new was re-indented);
    in several, or where new's indentation cannot be fitted: why not; nowhere: None."""
    lines = text.split("\n")
    o, n = old.split("\n"), new.split("\n")
    while o and not o[0].strip():                       # blank lines around old do not count (nor new's own)
        o.pop(0)
        if n and not n[0].strip():
            n.pop(0)
    while o and not o[-1].strip():
        o.pop()
        if n and not n[-1].strip():
            n.pop()
    if not o:
        return None
    want = [_squeeze(l) for l in o]
    have = [_squeeze(l) for l in lines]
    size = len(want)
    at = [i for i in range(len(have) - size + 1) if have[i] == want[0] and have[i:i + size] == want]
    if not at:
        return None
    if len(at) > 1:
        return (f"it matches {len(at)} places when whitespace is ignored (lines "
                f"{', '.join(str(i + 1) for i in at[:6])}). Add surrounding lines, copied exactly.")
    i = at[0]
    there = lines[i:i + size]
    fitted = _fit_indent(n, o, there)
    if fitted is None:
        exact, _ = clip("\n".join(there), 2000)
        return (f"it matches lines {i + 1}-{i + size} only when whitespace is ignored, and new_string's indentation "
                f"could not be fitted to the file's. Those lines exactly:\n{exact}")
    return "\n".join(lines[:i] + fitted + lines[i + size:]), i + 1, fitted != n


def _closest(text: str, old: str) -> str:
    """Where the file has the text most like old, as a hint: those lines exactly, to copy from. "" when nothing is
    close (or the file is too long to look through)."""
    lines = text.split("\n")
    o = old.strip("\n").split("\n")
    size = len(o)
    if len(lines) > 20000 or size > 300:
        return ""
    have = [_squeeze(l) for l in lines]
    sm = SequenceMatcher(autojunk=False)
    sm.set_seq2("\n".join(_squeeze(l) for l in o))
    best, at = 0.6, -1                                   # less alike than this is no help
    for i in range(max(1, len(lines) - size + 1)):
        sm.set_seq1("\n".join(have[i:i + size]))
        if sm.real_quick_ratio() <= best or sm.quick_ratio() <= best:
            continue
        r = sm.ratio()
        if r > best:
            best, at = r, i
    if at < 0:
        return ""
    last = min(len(lines), at + size)
    exact, _ = clip("\n".join(lines[at:last]), 2000)
    return f" The most similar text is at lines {at + 1}-{last}, exactly:\n{exact}\n"


async def edit_file(ctx: ToolContext, args: dict) -> ToolResult:
    rel = args.get("path") or ""
    old, new = args.get("old_string"), args.get("new_string")
    if not isinstance(old, str) or not isinstance(new, str) or old == "":
        return ToolResult("old_string and new_string are required (old_string not empty).", is_error=True)
    p = ctx.workdir.resolve(rel)
    if not p.is_file():
        return ToolResult(f"No such file: {rel}", is_error=True)
    raw = p.read_bytes().decode("utf-8", "replace")
    crlf = raw.count("\r\n") > raw.count("\n") - raw.count("\r\n")     # written back with the line ends it had
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    old, new = old.replace("\r\n", "\n"), new.replace("\r\n", "\n")
    count, how, line = text.count(old), "", 0
    plain = _unnumbered(old) if count == 0 else None
    if plain is not None:
        # a copy of read_file's lines with their numbers (local models do this): the text without them
        old, new, count = plain, _unnumbered(new, every=False) or "", text.count(plain)
        how = " (old_string had read_file's line numbers in it - they were left out)"
    if count == 0:
        loose = _loose_edit(text, old, new)
        if isinstance(loose, str):
            return ToolResult(f"old_string was not found in {rel} exactly: {loose}", is_error=True)
        if loose is None:
            hint = _closest(text, old) or " "
            return ToolResult(f"old_string was not found in {rel}.{hint}Copy the text exactly (read the file again "
                              "if it may have changed).", is_error=True)
        updated, line, fitted = loose
        how += " (found with whitespace ignored" + (", new_string indented the file's way" if fitted else "") + ")"
    else:
        if count > 1 and not args.get("replace_all"):
            return ToolResult(f"old_string occurs {count} times in {rel}. Add surrounding lines to make it unique, "
                              "or set replace_all.", is_error=True)
        updated = text.replace(old, new) if args.get("replace_all") else text.replace(old, new, 1)
        line = text[: text.index(old)].count("\n") + 1
    mark = copied_mark(new, text)
    if mark:
        return marked(mark, rel)
    ctx.workdir.write(rel, updated.replace("\n", "\r\n") if crlf else updated, "agent", "hub", f"agent edited {rel}")
    n = count if count and args.get("replace_all") else 1
    return ToolResult(f"Edited {rel}: replaced {n} occurrence(s), starting at line {line}.{how}")


async def search_files(ctx: ToolContext, args: dict) -> ToolResult:
    pattern = args.get("pattern") or ""
    try:
        rx = re.compile(pattern, 0 if args.get("case_sensitive") else re.IGNORECASE)
    except re.error as exc:
        return ToolResult(f"Not a valid regular expression: {exc}", is_error=True)
    glob = args.get("glob")
    base = ctx.workdir.resolve(args.get("path") or ".")
    limit = max(1, min(int(args.get("max_results") or 100), 500))
    hits = []
    for e in ctx.workdir.tree(ctx.workdir.rel(base)):
        if e["dir"] or e["size"] > 5_000_000:
            continue
        if glob and not Path(e["path"]).match(glob):
            continue
        p = ctx.workdir.resolve(e["path"])
        try:
            data = p.read_bytes()
        except OSError:
            continue
        if not is_text(data):
            continue
        for i, line in enumerate(data.decode("utf-8", "replace").splitlines(), 1):
            if rx.search(line):
                hits.append(f"{e['path']}:{i}: {line.strip()[:300]}")
                if len(hits) >= limit:
                    break
        if len(hits) >= limit:
            break
    if not hits:
        return ToolResult(f"No matches for /{pattern}/.")
    return ToolResult("\n".join(hits) + (f"\n(stopped at {limit} matches)" if len(hits) >= limit else ""))


async def move_file(ctx: ToolContext, args: dict) -> ToolResult:
    ctx.workdir.move(args.get("source") or "", args.get("destination") or "", "agent", "hub")
    return ToolResult(f"Moved {args.get('source')} to {args.get('destination')}.")


async def delete_file(ctx: ToolContext, args: dict) -> ToolResult:
    ctx.workdir.delete(args.get("path") or "", "agent", "hub")
    extra = " It can be restored from the project's history." if ctx.workdir.project_id else ""
    return ToolResult(f"Deleted {args.get('path')}.{extra}")


register(Tool("list_files", "List the files in the working folder (the project, or this chat's folder), "
              "as an indented tree with sizes.",
              obj({"path": {"type": "string", "description": "Sub-folder (default: the whole folder)"},
                   "pattern": {"type": "string", "description": "Glob for file names, e.g. *.py"},
                   "depth": {"type": "integer", "description": "How many levels deep (default 3)"}}),
              list_files, read_only=True))
register(Tool("read_file", "Read a file from the working folder. Text comes back with line numbers; "
              "images are shown to you; PDF, Word, Excel (every sheet, values and formulas) and PowerPoint (slides "
              "and notes) are read as text. Use offset/limit for long files (lines; pages for PDFs).",
              obj({"path": {"type": "string"}, "offset": {"type": "integer"}, "limit": {"type": "integer"}},
                  ["path"]), read_file, read_only=True))
register(Tool("write_file", "Create or overwrite a text file in the working folder, or with append: true add to "
              "its end. Prefer edit_file for changes to an existing file. Write a long file in parts: the first with "
              "write_file, each next one with append: true.",
              obj({"path": {"type": "string"},
                   "content": {"type": "string", "description": "The whole file - or with append, the next part"},
                   "append": {"type": "boolean", "description": "Add content at the end of the file (it is created "
                                                                "if it does not exist)"},
                   "replace": {"type": "boolean", "description": "Start over: replace a file you are writing in "
                                                                 "parts in this turn (without it, that is refused)"}},
                  ["path", "content"]), write_file))
register(Tool("edit_file", "Replace exact text in a file. old_string must match the file exactly (whitespace "
              "included) and be unique unless replace_all is true. Read the file first.",
              obj({"path": {"type": "string"}, "old_string": {"type": "string"}, "new_string": {"type": "string"},
                   "replace_all": {"type": "boolean"}}, ["path", "old_string", "new_string"]), edit_file))
register(Tool("search_files", "Search the text files of the working folder with a regular expression. "
              "Returns file:line: text.",
              obj({"pattern": {"type": "string"}, "path": {"type": "string"}, "glob": {"type": "string"},
                   "case_sensitive": {"type": "boolean"}, "max_results": {"type": "integer"}}, ["pattern"]),
              search_files, read_only=True))
register(Tool("move_file", "Move or rename a file or folder in the working folder.",
              obj({"source": {"type": "string"}, "destination": {"type": "string"}}, ["source", "destination"]),
              move_file))
register(Tool("delete_file", "Delete a file or folder in the working folder (project files stay in the history).",
              obj({"path": {"type": "string"}}, ["path"]), delete_file))
