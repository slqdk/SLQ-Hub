"""0.23.0 - skills for everyone: each person's own (private until shared with the family), made with the assistant
from their material (skill_propose and the card's Add skill), picked by hand for a chat or a project, in the format
Claude uses (import a .zip from Claude, download one for Claude)."""
from __future__ import annotations

import io
import json
import zipfile

import openpyxl
import pytest

from app.agent.skills import Skills, check_skill_md
from app.agent.tools.files import document_text

from .test_agent import finish, new_chat, say
from .test_audio import chat_dir
from .test_people import add_person, client_from, sign_in, turn, user_msg

ORIGIN = {"Origin": "http://hub"}
OWNER = None                                        # the owner's view: who=None
MARIA = {"id": 5, "role": "member"}
PER = {"id": 7, "role": "member"}


def md(name: str, desc: str = "", body: str = "") -> bytes:
    return (f"---\nname: {name}\ndescription: {desc or f'The {name} skill.'}\n---\n"
            f"{body or f'The {name} instructions.'}\n").encode()


def bundle(top: str, files: dict[str, bytes]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for rel, data in files.items():
            z.writestr(f"{top}/{rel}" if top else rel, data)
    return buf.getvalue()


def xlsx(rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- the store

def test_each_has_their_own_skills_and_sees_what_the_family_shares(tmp_path):
    s = Skills(tmp_path / "skills")
    assert [x["ref"] for x in s.visible(OWNER)] == ["hub/skill-creator"]          # the Hub's own, for everyone
    assert s.install(OWNER, md("haiku-writer"), "SKILL.md") == "haiku-writer"
    assert s.install(MARIA, bundle("Jule Frokost", {"SKILL.md": md("Jule Frokost"), "references/menu.md": b"Sild"}),
                     "jul.zip") == "@5/jule-frokost"                                 # a name the format takes
    assert s.install(PER, md("per-notes"), "SKILL.md") == "@7/per-notes"
    refs = lambda who: [x["ref"] for x in s.visible(who)]                            # noqa: E731
    # new skills are private: each sees their own and the Hub's
    assert refs(OWNER) == ["haiku-writer", "hub/skill-creator"]
    assert refs(MARIA) == ["@5/jule-frokost", "hub/skill-creator"]
    s.set_shared(OWNER, "haiku-writer", True)
    s.set_shared(MARIA, "@5/jule-frokost", True)
    assert refs(MARIA) == ["@5/jule-frokost", "haiku-writer", "hub/skill-creator"]
    assert refs(PER) == ["@7/per-notes", "haiku-writer", "@5/jule-frokost", "hub/skill-creator"]
    assert refs(OWNER) == ["haiku-writer", "@5/jule-frokost", "hub/skill-creator"]
    mine = {x["ref"]: x["mine"] for x in s.visible(MARIA)}
    assert mine == {"@5/jule-frokost": True, "haiku-writer": False, "hub/skill-creator": False}
    assert s.entry(PER, "@5/jule-frokost")["files"] == ["references/menu.md"]
    assert s.read_file(PER, "@5/jule-frokost", "references/menu.md") == b"Sild"
    # only the maker changes or deletes a skill; a private one cannot be read by anyone else
    for who, ref in ((MARIA, "haiku-writer"), (OWNER, "@5/jule-frokost"), (PER, "hub/skill-creator")):
        with pytest.raises(PermissionError):
            s.put_text(who, ref, "x")
        with pytest.raises(PermissionError):
            s.delete(who, ref)
    with pytest.raises(FileNotFoundError):
        s.get(MARIA, "@7/per-notes")
    assert not s.may_see(OWNER, "@7/per-notes") and not s.may_see(MARIA, "../x") and not s.may_see(MARIA, "@5/..")
    # taking it back
    s.set_shared(MARIA, "@5/jule-frokost", False)
    assert refs(PER) == ["@7/per-notes", "haiku-writer", "hub/skill-creator"]
    # a skill from before 0.23.0 (no .disabled.json entry) stays shared, as "switched on" meant
    old = tmp_path / "skills" / "old-skill"
    old.mkdir()
    (old / "SKILL.md").write_bytes(md("old-skill"))
    assert s.may_see(MARIA, "old-skill")
    # a removed person's skills go with them
    s.forget_person(5)
    assert "@5/jule-frokost" not in refs(OWNER) and not (tmp_path / "skills" / "@5").exists()


def test_categories_are_kept_by_each_skills_owner(tmp_path):
    """0.41.0: the owner of a skill puts it into a category; who sees it - the owner, or the family when shared -
    sees it there. Renaming a category moves all of one's own skills in it; a deleted skill leaves its category."""
    s = Skills(tmp_path / "skills")
    for name in ("tc-plc", "tc-motion", "pdf-guide"):
        s.install(OWNER, md(name), "SKILL.md")
    s.install(MARIA, md("jule"), "SKILL.md")
    cat = lambda who: {x["ref"]: x["category"] for x in s.visible(who)}             # noqa: E731
    assert s.set_category(OWNER, "tc-plc", "  TwinCAT   PLC ") == "TwinCAT PLC"
    s.set_category(OWNER, "tc-motion", "TwinCAT PLC")
    s.set_category(OWNER, "pdf-guide", "Documents")
    s.set_category(MARIA, "@5/jule", "Food")
    assert cat(OWNER) == {"tc-plc": "TwinCAT PLC", "tc-motion": "TwinCAT PLC", "pdf-guide": "Documents",
                          "hub/skill-creator": ""}
    s.set_shared(OWNER, "tc-plc", True)
    s.set_shared(OWNER, "pdf-guide", False)
    assert cat(MARIA)["tc-plc"] == "TwinCAT PLC" and cat(MARIA)["@5/jule"] == "Food" and "pdf-guide" not in cat(MARIA)
    for who, ref in ((MARIA, "tc-plc"), (OWNER, "@5/jule"), (OWNER, "hub/skill-creator")):
        with pytest.raises(PermissionError):
            s.set_category(who, ref, "Mine now")
    with pytest.raises(ValueError):
        s.set_category(OWNER, "tc-plc", "x" * 41)
    assert s.rename_category(OWNER, "twincat plc", "TwinCAT") == 2
    assert s.rename_category(MARIA, "TwinCAT", "Hers") == 0                          # only one's own skills move
    assert cat(OWNER)["tc-motion"] == "TwinCAT" and cat(MARIA)["tc-plc"] == "TwinCAT"
    assert s.set_category(OWNER, "pdf-guide", "") == "" and cat(OWNER)["pdf-guide"] == ""
    s.delete(OWNER, "tc-motion")
    s.install(OWNER, md("tc-motion"), "SKILL.md")
    assert cat(OWNER)["tc-motion"] == ""                                            # made again: in none
    s.install(OWNER, md("tc-plc", "Changed."), "SKILL.md")                           # replaced: keeps its category
    assert cat(OWNER)["tc-plc"] == "TwinCAT"


def test_the_format_is_claudes_both_ways(tmp_path):
    s = Skills(tmp_path / "skills")
    # a skill downloaded from Claude: one folder with SKILL.md at its top (and macOS litter)
    ref = s.install(OWNER, bundle("pdf-tools", {"SKILL.md": md("pdf-tools", "Fill PDF forms. Use for PDFs."),
                                                "scripts/fill.py": b"print(1)", "../evil.txt": b"x"}),
                    "pdf-tools.zip")
    assert ref == "pdf-tools" and s.entry(OWNER, ref)["files"] == ["scripts/fill.py"]
    assert not (tmp_path / "evil.txt").exists() and not (tmp_path / "skills" / "evil.txt").exists()
    name, data = s.zip(OWNER, ref)
    z = zipfile.ZipFile(io.BytesIO(data))
    assert name == "pdf-tools" and sorted(z.namelist()) == ["pdf-tools/SKILL.md", "pdf-tools/scripts/fill.py"]
    # the same name again replaces it - and keeps whether it was shared
    s.set_shared(OWNER, ref, True)
    s.install(OWNER, md("pdf-tools", "New text."), "SKILL.md")
    assert s.entry(OWNER, ref)["description"] == "New text." and s.entry(OWNER, ref)["shared"] is True
    for bad, words in ((b"PK\x03\x04junk", "not a zip"), (bundle("x", {"notes.md": b"hi"}), "no SKILL.md")):
        with pytest.raises(ValueError, match=words):
            s.install(OWNER, bad, "x.zip")
    # metadata maps and folded descriptions are read as agentskills.io writes them
    meta, problems = check_skill_md("---\nname: data-tool\ndescription: >\n  Reads data.\n  Use for data.\nmetadata:\n"
                                    "  author: soren\n  version: \"1.0\"\n---\nDo it.\n")
    assert problems == [] and meta["description"] == "Reads data. Use for data." and \
        meta["metadata"] == {"author": "soren", "version": "1.0"}
    for text, words in (("no front matter", "must start with front matter"),
                        ("---\nname: Data Tool\ndescription: x\n---\nDo.", "lowercase letters"),
                        ("---\nname: my--tool\ndescription: x\n---\nDo.", "single hyphens"),
                        ("---\nname: claude-helper\ndescription: x\n---\nDo.", "claude, anthropic"),
                        ("---\nname: tool\n---\nDo.", "no description"),
                        ("---\nname: tool\ndescription: " + "x" * 1025 + "\n---\nDo.", "at most 1024"),
                        ("---\nname: tool\ndescription: a <b>bold</b> one\n---\nDo.", "XML tags"),
                        ("---\nname: tool\ndescription: x\n---\n", "no instructions")):
        assert any(words in p for p in check_skill_md(text)[1]), (text, check_skill_md(text)[1])


def test_excel_and_powerpoint_are_read_as_text():
    text = document_text("budget.xlsx", xlsx([["Måned", "Beløb"], ["januar", 1200], ["Sum", "=SUM(B2:B2)"]]))[1]
    assert "--- sheet Sheet ---" in text and "1: Måned | Beløb" in text and "2: januar | 1200" in text and \
        "3: Sum | =SUM(B2:B2)" in text
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Nødstop"
    slide.notes_slide.notes_text_frame.text = "Vis knappen"
    buf = io.BytesIO()
    deck.save(buf)
    text = document_text("kursus.pptx", buf.getvalue())[1]
    assert text == "--- slide 1 ---\nNødstop\nNotes: Vis knappen"
    assert document_text("a.txt", b"x") is None


# ---------------------------------------------------------------- picked by hand

def tools_of(req: dict) -> set[str]:
    return {t["function"]["name"] for t in req.get("tools") or []}


async def test_only_the_skills_chosen_for_a_chat_reach_the_model(signed_in, app, fakes):
    for name in ("haiku-writer", "plc-notes"):
        files = {"SKILL.md": md(name, body=f"The {name} instructions.")}
        if name == "plc-notes":
            files["references/codes.xlsx"] = xlsx([["Kode", "Betydning"], ["E1", "Nødstop"]])
            files["scripts/check.py"] = b"print('ok')"
        r = await signed_in.post("/hub/api/skills", files={"file": (f"{name}.zip", bundle(name, files),
                                                                    "application/zip")}, headers=ORIGIN)
        assert r.status_code == 200, r.text
    c = await new_chat(signed_in)
    fakes.llama["replies"] = [{"content": "Hi"}, {"content": "T"}]
    await finish(app, await say(signed_in, c["id"], "hi"))
    req = fakes.llama["requests"][0]
    assert "# Skills" not in req["messages"][0]["content"] and "skill" not in tools_of(req)  # none chosen: none
    await signed_in.patch(f"/hub/api/chats/{c['id']}", json={"settings": {"skills": ["plc-notes"]}})
    fakes.llama["requests"].clear()
    fakes.llama["replies"] = [
        {"tool_calls": [{"name": "skill", "arguments": {"name": "plc-notes", "file": "references/codes.xlsx"}}]},
        {"tool_calls": [{"name": "skill", "arguments": {"name": "haiku-writer"}}]},
        {"content": "E1 is the emergency stop."}]
    run = await finish(app, await say(signed_in, c["id"], "what is E1?"))
    system = fakes.llama["requests"][0]["messages"][0]["content"]
    assert "## Skill: plc-notes (plc-notes)\nThe plc-notes instructions." in system
    assert ".hub/skills/plc-notes/" in system and "haiku-writer" not in system
    results = [e["data"] for e in run.events if e["type"] == "tool_result"]
    assert "2: E1 | Nødstop" in results[0]["text"]
    assert not results[1]["ok"] and "No skill 'haiku-writer' in this chat" in results[1]["text"]
    copy = chat_dir(app, c["id"]) / ".hub" / "skills" / "plc-notes"
    assert (copy / "scripts" / "check.py").read_bytes() == b"print('ok')"            # for code
    files = await signed_in.get(f"/hub/api/folders/chat/{c['id']}/tree")
    assert files.status_code == 200 and ".hub" not in json.dumps(files.json())      # not among the chat's files
    # a project's skills go with its chats; a skill no longer chosen leaves the folder
    p = (await signed_in.post("/hub/api/projects", json={"name": "Haiku"})).json()
    await signed_in.patch(f"/hub/api/projects/{p['id']}", json={"settings": {"skills": ["haiku-writer"]}})
    pc = await new_chat(signed_in, project_id=p["id"])
    fakes.llama["requests"].clear()
    fakes.llama["replies"] = [{"content": "Autumn moon"}, {"content": "T"}]
    await finish(app, await say(signed_in, pc["id"], "a haiku"))
    assert "## Skill: haiku-writer (haiku-writer)" in fakes.llama["requests"][0]["messages"][0]["content"]
    await signed_in.patch(f"/hub/api/chats/{c['id']}", json={"settings": {"skills": []}})
    fakes.llama["replies"] = [{"content": "ok"}]
    await finish(app, await say(signed_in, c["id"], "again"))
    assert not copy.exists()


async def test_a_long_skill_is_loaded_rather_than_put_in_the_prompt(signed_in, app, fakes):
    body = "Step.\n" * 8000                                                            # 48,000 characters
    await signed_in.post("/hub/api/skills", files={"file": ("SKILL.md", md("big-manual", body=body), "text/markdown")},
                         headers=ORIGIN)
    c = await new_chat(signed_in, settings={"skills": ["big-manual"]})
    fakes.llama["replies"] = [{"content": "ok"}, {"content": "T"}]
    await finish(app, await say(signed_in, c["id"], "go"))
    system = fakes.llama["requests"][0]["messages"][0]["content"]
    assert 'Its instructions are long: load them with skill(name="big-manual") before you start.' in system
    assert "Step.\nStep." not in system and "skill" in tools_of(fakes.llama["requests"][0])
    # a skill of SKILL.md alone is in the working folder too (the skill maker improves it from there)
    assert "Its folder, for code or to improve it: .hub/skills/big-manual/" in system
    assert (chat_dir(app, c["id"]) / ".hub" / "skills" / "big-manual" / "SKILL.md").read_bytes().startswith(b"---\nname: big-manual")


# ---------------------------------------------------------------- made with the assistant

SKILL = ("---\nname: julefrokost-invitation\ndescription: Writes the firm's Christmas lunch invitation in Danish with "
         "the date, place and sign-up date. Use for julefrokost invitations.\n---\n# Invitation\n1. Four points.\n"
         "2. Read references/rules.md.\n")


async def test_a_skill_made_with_the_assistant_is_checked_and_added_from_its_card(signed_in, app, fakes):
    c = await new_chat(signed_in, settings={"skills": ["hub/skill-creator"]})
    folder = chat_dir(app, c["id"])
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "skabelon.xlsx").write_bytes(xlsx([["Navn", "Tilmeldt"]]))
    bad = SKILL.replace("name: julefrokost-invitation", "name: Julefrokost Invitation")
    fakes.llama["replies"] = [
        {"tool_calls": [{"name": "write_file", "arguments": {"path": "skills/julefrokost-invitation/SKILL.md",
                                                             "content": bad}}]},
        {"tool_calls": [{"name": "skill_propose", "arguments": {"path": "skills/julefrokost-invitation"}}]},
        {"tool_calls": [{"name": "write_file", "arguments": {"path": "skills/julefrokost-invitation/SKILL.md",
                                                             "content": SKILL, "replace": True}}]},
        {"tool_calls": [{"name": "write_file", "arguments": {"path": "skills/julefrokost-invitation/references/"
                                                                     "rules.md", "content": "No 'hyggelig'."}}]},
        {"tool_calls": [{"name": "skill_propose", "arguments": {"path": "skills/julefrokost-invitation",
                                                                "include": ["skabelon.xlsx"]}}]},
        {"content": "Press Add skill on the card."}, {"content": "Skill"}]
    run = await finish(app, await say(signed_in, c["id"], "Make a skill for our julefrokost invitations"))
    assert run.status == "done", run.events[-1]
    system = fakes.llama["requests"][0]["messages"][0]["content"]
    assert "## Skill: skill-creator (hub/skill-creator)" in system and "# Making a skill" in system
    assert "skill_propose" in tools_of(fakes.llama["requests"][0])
    results = [e["data"] for e in run.events if e["type"] == "tool_result"]
    assert not results[1]["ok"] and "The skill is not ready: name 'Julefrokost Invitation' must be" in results[1]["text"]
    ok = results[4]
    assert ok["ok"] and "julefrokost-invitation is ready (3 files" in ok["text"] and "Add skill" in ok["text"]
    card = ok["display"]["skill"]
    assert card["name"] == "julefrokost-invitation" and card["replaces"] is False
    assert card["files"] == ["assets/skabelon.xlsx", "references/rules.md"] and "text" not in card
    assert (folder / "skills" / "julefrokost-invitation" / "assets" / "skabelon.xlsx").is_file()
    # not added until the user presses Add skill
    refs = [s["ref"] for s in (await signed_in.get("/hub/api/skills")).json()]
    assert "julefrokost-invitation" not in refs
    d = (await signed_in.get(f"/hub/api/skills/drafts/{card['draft']}")).json()
    assert d["text"].startswith("---\nname: julefrokost-invitation") and d["added"] is None
    z = zipfile.ZipFile(io.BytesIO((await signed_in.get(f"/hub/api/skills/drafts/{card['draft']}/zip")).content))
    assert "julefrokost-invitation/references/rules.md" in z.namelist()
    r = await signed_in.post(f"/hub/api/skills/drafts/{card['draft']}/add", headers=ORIGIN)
    assert r.json()["ref"] == "julefrokost-invitation" and r.json()["skill"]["mine"] is True
    assert r.json()["skill"]["shared"] is False                                       # private until shared
    assert (await signed_in.get(f"/hub/api/skills/drafts/{card['draft']}")).json()["added"]
    # improving it: the same name proposed again replaces it when added
    (folder / "skills" / "julefrokost-invitation" / "SKILL.md").write_text(SKILL + "3. Sign with the firm's name.\n")
    store = app.state.hub.extras["skills"]
    again = store.propose(None, folder / "skills" / "julefrokost-invitation")
    assert again["replaces"] is True
    await signed_in.post(f"/hub/api/skills/drafts/{again['draft']}/add", headers=ORIGIN)
    text = (await signed_in.get("/hub/api/skills/one", params={"ref": "julefrokost-invitation"})).json()["text"]
    assert "3. Sign with the firm's name." in text
    assert (await signed_in.get("/hub/api/skills/drafts/0123456789abcdef")).status_code == 404


# ---------------------------------------------------------------- people

def test_a_skill_can_name_its_category(tmp_path):
    """0.44.0: metadata: category: in the front matter puts an imported skill there - once; the owner decides after."""
    s = Skills(tmp_path / "skills")
    text = b"---\nname: billede-med-dyr\ndescription: Dyr.\nmetadata:\n  category: Billeder\n---\nDo it.\n"
    s.install(OWNER, text, "SKILL.md")
    assert s.entry(OWNER, "billede-med-dyr")["category"] == "Billeder"
    s.set_category(OWNER, "billede-med-dyr", "Mine billeder")
    s.install(OWNER, text, "SKILL.md")                                              # replaced: the owner's stands
    assert s.entry(OWNER, "billede-med-dyr")["category"] == "Mine billeder"
    s.install(MARIA, text, "SKILL.md")                                              # each owner gets their own
    assert s.entry(MARIA, "@5/billede-med-dyr")["category"] == "Billeder"


async def test_categories_over_the_api(signed_in, app, fakes):
    r = await signed_in.post("/hub/api/skills", files={"file": ("SKILL.md", md("pdf-guide"), "text/markdown")})
    assert r.status_code == 200, r.text
    r = await signed_in.patch("/hub/api/skills/one", params={"ref": "pdf-guide"}, json={"category": "Documents"})
    assert r.status_code == 200 and r.json()["category"] == "Documents"
    listed = {x["ref"]: x for x in (await signed_in.get("/hub/api/skills")).json()}
    assert listed["pdf-guide"]["category"] == "Documents" and listed["pdf-guide"]["shared"] is False
    r = await signed_in.patch("/hub/api/skills/one", params={"ref": "pdf-guide"}, json={"shared": True})
    listed = {x["ref"]: x for x in (await signed_in.get("/hub/api/skills")).json()}
    assert listed["pdf-guide"]["category"] == "Documents" and listed["pdf-guide"]["shared"] is True
    r = await signed_in.post("/hub/api/skills/categories/rename", json={"old": "documents", "new": "Docs"})
    assert r.json() == {"moved": 1}
    r = await signed_in.patch("/hub/api/skills/one", params={"ref": "hub/skill-creator"}, json={"category": "X"})
    assert r.status_code == 403
    r = await signed_in.patch("/hub/api/skills/one", params={"ref": "pdf-guide"}, json={"category": "y" * 41})
    assert r.status_code == 400


async def test_people_make_share_and_use_their_own_skills(signed_in, app, fakes):
    for name in ("haiku-writer", "owner-private"):
        await signed_in.post("/hub/api/skills", files={"file": ("SKILL.md", md(name), "text/markdown")}, headers=ORIGIN)
    await signed_in.patch("/hub/api/skills/one", params={"ref": "haiku-writer"}, json={"shared": True}, headers=ORIGIN)
    maria = await add_person(signed_in)
    mine = f"@{maria['id']}/kage-opskrifter"
    async with client_from(app) as phone:
        await sign_in(phone)
        r = await phone.post("/hub/api/skills", files={"file": ("SKILL.md", md("kage-opskrifter"), "text/markdown")},
                             headers=ORIGIN)
        assert r.json()["ref"] == mine
        seen = {s["ref"]: s for s in (await phone.get("/hub/api/skills")).json()}
        assert list(seen) == [mine, "haiku-writer", "hub/skill-creator"]
        assert seen["haiku-writer"]["by"] == "Soren" and seen[mine]["mine"] and not seen[mine]["shared"]
        assert (await phone.get("/hub/api/skills/one", params={"ref": "owner-private"})).status_code == 404
        assert (await phone.delete("/hub/api/skills/one", params={"ref": "haiku-writer"},
                                   headers=ORIGIN)).status_code == 403
        assert (await phone.get("/hub/api/skills/anything")).status_code == 403          # not a skills route
        # a turn with the skills she chose: hers and the shared one, never the owner's private one
        fakes.llama["replies"] = [
            {"tool_calls": [{"name": "write_file", "arguments": {"path": "skills/strik-moenstre/SKILL.md",
                                                                 "content": md("strik-moenstre").decode()}}]},
            {"tool_calls": [{"name": "skill_propose", "arguments": {"path": "skills/strik-moenstre"}}]},
            {"content": "Done."}]
        res = await turn(app, phone, {"title": "S", "settings": {"skills": [mine, "haiku-writer", "owner-private"]},
                                      "messages": [user_msg("Make a knitting skill")]})
        system = fakes.llama["requests"][0]["messages"][0]["content"]
        assert f"## Skill: kage-opskrifter ({mine})" in system and "## Skill: haiku-writer" in system
        assert "owner-private" not in system
        card = next(b["display"]["skill"] for m in res["messages"] for b in m["content"]
                    if b["type"] == "tool_result" and (b.get("display") or {}).get("skill"))
        # her draft is hers: the owner cannot add it
        assert (await signed_in.post(f"/hub/api/skills/drafts/{card['draft']}/add", headers=ORIGIN)).status_code == 404
        r = await phone.post(f"/hub/api/skills/drafts/{card['draft']}/add", headers=ORIGIN)
        assert r.json()["ref"] == f"@{maria['id']}/strik-moenstre"
        owner_sees = [s["ref"] for s in (await signed_in.get("/hub/api/skills")).json()]
        assert mine not in owner_sees                                                   # private until she shares
        await phone.patch("/hub/api/skills/one", params={"ref": mine}, json={"shared": True}, headers=ORIGIN)
        shared = {s["ref"]: s for s in (await signed_in.get("/hub/api/skills")).json()}
        assert shared[mine]["by"] == "Maria" and not shared[mine]["mine"]
        name, = [n for n in zipfile.ZipFile(io.BytesIO((await phone.get(
            "/hub/api/skills/zip", params={"ref": mine})).content)).namelist()]
        assert name == "kage-opskrifter/SKILL.md"
    await signed_in.delete(f"/hub/api/people/{maria['id']}", headers=ORIGIN)
    assert mine not in [s["ref"] for s in (await signed_in.get("/hub/api/skills")).json()]


def test_internal_is_the_admins_alone(tmp_path):
    """0.47.0, Søren: "internal mappen skal ALTID kun være tilgængelig for admin". The owner's skills in internal
    are not listed, read or used for anyone else - shared before or not - and a person cannot make an internal of
    their own (one named so before is no category)."""
    s = Skills(tmp_path / "skills")
    for name in ("secret-plc", "family-pics"):
        s.install(OWNER, md(name), "SKILL.md")
        s.set_shared(OWNER, name, True)
    assert s.set_category(OWNER, "secret-plc", " Internal ") == "Internal"
    refs = lambda who: [x["ref"] for x in s.visible(who)]                           # noqa: E731
    assert "secret-plc" in refs(OWNER) and "secret-plc" not in refs(MARIA) and "family-pics" in refs(MARIA)
    mine = s.entry(OWNER, "secret-plc")
    assert mine["admin_only"] is True and mine["shared"] is False and mine["category"] == "Internal"
    assert not s.may_see(MARIA, "secret-plc") and s.chosen(["secret-plc", "family-pics"], MARIA)[0]["ref"] == \
        "family-pics" and len(s.chosen(["secret-plc"], MARIA)) == 0
    for read in (lambda: s.get(MARIA, "secret-plc"), lambda: s.path(MARIA, "secret-plc"),
                 lambda: s.entry(MARIA, "secret-plc"), lambda: s.zip(MARIA, "secret-plc")):
        with pytest.raises(FileNotFoundError):
            read()
    with pytest.raises(ValueError):
        s.set_shared(OWNER, "secret-plc", True)
    s.set_category(OWNER, "secret-plc", "TwinCAT")                                  # out again: still private
    assert s.entry(OWNER, "secret-plc")["shared"] is False and "secret-plc" not in refs(MARIA)
    s.set_category(OWNER, "family-pics", "Pictures")
    assert s.rename_category(OWNER, "pictures", "internal") == 1 and "family-pics" not in refs(MARIA)
    # a person: no internal of their own - set, renamed into or brought along by an imported skill
    s.install(MARIA, md("jule"), "SKILL.md")
    s.set_category(MARIA, "@5/jule", "Food")
    for change in (lambda: s.set_category(MARIA, "@5/jule", "INTERNAL"),
                   lambda: s.rename_category(MARIA, "Food", "internal")):
        with pytest.raises(ValueError):
            change()
    s.install(MARIA, b"---\nname: kage\ndescription: Kage.\nmetadata:\n  category: internal\n---\nBag.\n", "SKILL.md")
    assert s.entry(MARIA, "@5/kage")["category"] == ""
    (tmp_path / "skills" / "@5" / ".categories.json").write_text(json.dumps({"jule": "internal", "kage": "Food"}))
    assert {x["ref"]: x["category"] for x in s.visible(MARIA) if x["mine"]} == {"@5/jule": "", "@5/kage": "Food"}


async def test_internal_is_not_seen_from_a_persons_phone(signed_in, app, fakes):
    await signed_in.post("/hub/api/skills", files={"file": ("SKILL.md", md("admin-tools"), "text/markdown")},
                         headers=ORIGIN)
    await signed_in.patch("/hub/api/skills/one", params={"ref": "admin-tools"}, json={"shared": True}, headers=ORIGIN)
    r = await signed_in.patch("/hub/api/skills/one", params={"ref": "admin-tools"}, json={"category": "internal"},
                              headers=ORIGIN)
    assert r.status_code == 200
    r = await signed_in.patch("/hub/api/skills/one", params={"ref": "admin-tools"}, json={"shared": True},
                              headers=ORIGIN)
    assert r.status_code == 400 and "internal" in r.json()["detail"]
    own = {x["ref"]: x for x in (await signed_in.get("/hub/api/skills")).json()}
    assert own["admin-tools"]["admin_only"] is True
    await add_person(signed_in)
    async with client_from(app) as phone:
        await sign_in(phone)
        assert "admin-tools" not in {x["ref"] for x in (await phone.get("/hub/api/skills")).json()}
        assert (await phone.get("/hub/api/skills/one", params={"ref": "admin-tools"})).status_code == 404
        assert (await phone.get("/hub/api/skills/zip", params={"ref": "admin-tools"})).status_code == 404
        fakes.llama["replies"] = [{"content": "Done."}]
        await turn(app, phone, {"title": "S", "settings": {"skills": ["admin-tools"]},
                                "messages": [user_msg("Hello")]})
        assert "admin-tools" not in fakes.llama["requests"][0]["messages"][0]["content"]
