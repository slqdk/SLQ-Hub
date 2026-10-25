"""The dev server's fake model on Settings -> Tests: answers each test case as a good model would, with the real
tools (so the sandbox's checks look at real files: a chart, a Word file, a C# program, a cut-out from the fake
ComfyUI) - for trying the Tests page and the checks end to end without a GPU. Not used by the Hub itself."""
from __future__ import annotations

import json
import re

from app.evals import fixtures

MONTHS = json.dumps(list(fixtures.MONTHS_DA))

SALES = f"""python3 - <<'EOF'
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
df = pd.read_csv(".hub/attachments/salg_2025.csv", sep=";", decimal=",", thousands=".")
df["dato"] = pd.to_datetime(df["dato"], format="%d-%m-%Y")
df["oms"] = df["antal"] * df["stykpris"]
q2 = df[df["dato"].dt.month.isin([4, 5, 6])].groupby("region")["oms"].sum().sort_values(ascending=False)
m = df.groupby(df["dato"].dt.month)["oms"].sum()
names = {MONTHS}
ax = m.plot.bar(color="#2f6f9f", figsize=(9, 4.5))
ax.set_xticklabels([n[:3] for n in names], rotation=0)
ax.set_ylabel("kr.")
ax.set_title("Omsætning pr. måned 2025")
plt.tight_layout()
plt.savefig("omsaetning.png", dpi=100)
amount = f"{{q2.iloc[0]:,.2f}}".replace(",", "X").replace(".", ",").replace("X", ".")
print(f"REGION: {{q2.index[0]}}; {{amount}} kr.")
print(f"MÅNED: {{names[m.idxmax() - 1]}}")
EOF"""

DOCX = """python3 - <<'EOF'
import docx
d = docx.Document()
d.add_heading("Månedsrapport oktober 2026", 0)
d.add_heading("Produktion", 1)
d.add_paragraph("Produktionen lå over plan i oktober, med færre stop end i september.")
t = d.add_table(rows=4, cols=3)
t.style = "Table Grid"
for r, row in enumerate([("Linje", "Produceret", "Kassation %"), ("Linje 1", "12.400", "1,8"),
                         ("Linje 2", "9.850", "2,4"), ("Linje 3", "11.200", "1,1")]):
    for c, v in enumerate(row):
        t.cell(r, c).text = v
d.add_heading("Kvalitet", 1)
d.add_paragraph("Kassationen faldt på alle tre linjer efter justeringen af doseringen.")
d.add_heading("Plan for november", 1)
for x in ("Service af linje 2 i uge 46", "Oplæring af to nye operatører", "Test af ny emballage"):
    d.add_paragraph(x, style="List Bullet")
d.save("rapport.docx")
print("rapport.docx gemt")
EOF"""

XLSX = """python3 - <<'EOF'
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(["Måned", "Indtægter", "Udgifter", "Resultat"])
names = %s
for i, m in enumerate(names, start=2):
    ws.append([m, 80000 + 1500 * i, 62000 + 900 * i, f"=B{i}-C{i}"])
ws.append(["Total", "=SUM(B2:B13)", "=SUM(C2:C13)", "=SUM(D2:D13)"])
wb.save("budget.xlsx")
print("budget.xlsx gemt")
EOF""" % MONTHS

ROMAN = '''"""Roman numerals for 1-3999."""
import re

_VALUES = [(1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
           (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
_VALID = re.compile(r"^M{0,3}(CM|CD|D?C{0,3})(XC|XL|L?X{0,3})(IX|IV|V?I{0,3})$")


def to_roman(n: int) -> str:
    if not isinstance(n, int) or isinstance(n, bool) or not 1 <= n <= 3999:
        raise ValueError(f"not between 1 and 3999: {n!r}")
    out = []
    for value, sym in _VALUES:
        while n >= value:
            out.append(sym)
            n -= value
    return "".join(out)


def from_roman(s: str) -> int:
    if not isinstance(s, str) or not s or not _VALID.match(s):
        raise ValueError(f"not a Roman numeral: {s!r}")
    total, i = 0, 0
    for value, sym in _VALUES:
        while s.startswith(sym, i):
            total += value
            i += len(sym)
    return total
'''
ROMAN_TESTS = '''import unittest

from roman import from_roman, to_roman


class RomanTests(unittest.TestCase):
    def test_examples(self):
        self.assertEqual(to_roman(1994), "MCMXCIV")
        self.assertEqual(from_roman("MMXXVI"), 2026)

    def test_round_trip(self):
        for n in range(1, 4000):
            self.assertEqual(from_roman(to_roman(n)), n)

    def test_bad(self):
        for s in ("", "IIII", "VX", "IC", "MMMM"):
            with self.assertRaises(ValueError):
                from_roman(s)
        with self.assertRaises(ValueError):
            to_roman(0)


if __name__ == "__main__":
    unittest.main()
'''

FIXED = fixtures.INVENTORY.replace("self.items[sku] = qty\n", "self.items[sku] = self.items.get(sku, 0) + qty\n") \
    .replace("self.items[sku] = have - qty\n", "self.items[sku] = have - qty\n        if not self.items[sku]:\n"
                                               "            del self.items[sku]\n") \
    .replace("if qty < threshold", "if qty <= threshold").replace("return sku, int(qty)", "return sku.strip(), int(qty)")

CSPROJ = """<Project Sdk="Microsoft.NET.Sdk">
  <PropertyGroup>
    <OutputType>Exe</OutputType>
    <TargetFramework>net10.0</TargetFramework>
    <ImplicitUsings>enable</ImplicitUsings>
    <Nullable>enable</Nullable>
  </PropertyGroup>
</Project>
"""
PROGRAM = """using System.Text.RegularExpressions;

if (args.Length < 1)
{
    Console.Error.WriteLine("Brug: ordtaeller <fil>");
    return 1;
}
var counts = new Dictionary<string, int>();
foreach (Match m in Regex.Matches(File.ReadAllText(args[0]), @"\\p{L}+"))
{
    var w = m.Value.ToLowerInvariant();
    counts[w] = counts.GetValueOrDefault(w) + 1;
}
foreach (var kv in counts.OrderByDescending(kv => kv.Value).ThenBy(kv => kv.Key, StringComparer.Ordinal).Take(5))
    Console.WriteLine($"{kv.Key}: {kv.Value}");
return 0;
"""

PAGE = """<!doctype html>
<html lang="da"><head><meta charset="utf-8"><title>Celsius og Fahrenheit</title>
<style>body{font-family:system-ui;padding:2rem;max-width:28rem;margin:auto}label{display:block;margin:.8rem 0}
input{font-size:1.2rem;width:100%;padding:.4rem}button{margin-top:1rem;padding:.5rem 1rem}</style></head>
<body><h1>Omregner</h1>
<label>Celsius <input id="c" type="number" step="any"></label>
<label>Fahrenheit <input id="f" type="number" step="any"></label>
<button id="reset">Nulstil</button>
<script>
const c = document.getElementById("c"), f = document.getElementById("f");
const round = (x) => Math.round(x * 100) / 100;
c.addEventListener("input", () => { f.value = c.value === "" ? "" : round(c.value * 9 / 5 + 32); });
f.addEventListener("input", () => { c.value = f.value === "" ? "" : round((f.value - 32) * 5 / 9); });
document.getElementById("reset").addEventListener("click", () => { c.value = ""; f.value = ""; });
</script></body></html>"""

FILES = """python3 - <<'EOF'
import os, random
os.makedirs("maalinger", exist_ok=True)
rnd = random.Random(5)
with open("maalinger/data.csv", "w") as f:
    f.write("time,temperatur\\n")
    for h in range(24):
        f.write(f"{h},{rnd.uniform(15, 25):.1f}\\n")
open("maalinger/analyse.py", "w").write('''import csv
with open("data.csv") as f:
    temps = [float(r["temperatur"]) for r in csv.DictReader(f)]
with open("resultat.md", "w") as out:
    out.write(f"# Resultat\\\\n\\\\n- Middel: {sum(temps) / len(temps):.2f}\\\\n- Minimum: {min(temps):.2f}\\\\n"
              f"- Maksimum: {max(temps):.2f}\\\\n")
''')
EOF
cd maalinger && python3 analyse.py && cat resultat.md"""


def _last_tool(msgs: list[dict]) -> str:
    for m in reversed(msgs):
        if m["role"] == "tool":
            return m["content"] if isinstance(m["content"], str) else json.dumps(m["content"])
    return ""


def _call(name: str, **arguments) -> dict:
    return {"tool_calls": [{"name": name, "arguments": arguments}]}


def answer(text: str, msgs: list[dict], done: int) -> dict | None:
    """The scripted turn for a test case's message (None: not a test case)."""
    tool = _last_tool(msgs)
    if "Fire personer - Anna, Bent" in text:
        return {"content": "Hunden bor i hus 2, og det blå hus ligger lige til venstre for det røde, hvor fisken bor. "
                           "Prøver man mulighederne, passer kun: 1 Bent gul kat, 2 Anna blå hund, 3 Dennis rød fisk, "
                           "4 Cecilie grøn papegøje.\n\nSVAR: Dennis, gul"}
    if "A tank is filled by two pumps" in text:
        return {"content": "From 08:00 to 09:00 A and B fill 1/6 + 1/4 = 5/12. From 09:00 to 10:00 the net rate is "
                           "1/6 + 1/4 - 1/12 = 1/3, so the tank is 3/4 full at 10:00. After that A - C = 1/12 per hour: "
                           "the last 1/4 takes 3 hours.\n\nANSWER: 13:00"}
    if "Udtræk oplysningerne fra e-mailen" in text:
        from app.evals.suite import ORDER_WANT
        return {"content": "```json\n" + json.dumps(ORDER_WANT, ensure_ascii=False, indent=2) + "\n```"}
    if "invitation til julefrokosten" in text:
        return {"content": "Kære kolleger\n\nVi holder julefrokost, og I er alle inviteret:\n\n"
                           "- Fredag den 4. december fra kl. 17\n- Vi mødes i kantinen\n"
                           "- Tilmelding senest den 20. november\n- Tag godt humør og en lille gave med\n\n"
                           "Venlig hilsen\nFestudvalget"}
    if "salgsdata for 2025" in text:
        if done == 0:
            return _call("bash", command=SALES)
        lines = [x.strip() for x in tool.splitlines() if x.strip().startswith(("REGION:", "MÅNED:"))]
        return {"content": "Jeg læste filen med dansk talformat (semikolon, decimalkomma) og regnede antal × stykpris "
                           "sammen. Diagrammet er gemt som omsaetning.png.\n\n" + "\n".join(lines)}
    if "driftsmanualen for pumpestationen" in text:
        if done == 0:
            return _call("read_file", path=".hub/attachments/driftsmanual.pdf")
        return {"content": "Side 9 angiver 2.150 driftstimer for P3, men ændringsloggen på side 11 sænker intervallet "
                           "til 1.900 driftstimer fra 1. januar 2026.\n\nSVAR: 1.900 driftstimer"}
    if "serverloggen fra august til oktober" in text:
        if done == 0:
            return _call("bash", command="grep 'NAS-2 ERROR backup FAILED' .hub/attachments/server.log")
        sept = [x for x in tool.splitlines() if re.match(r"^2026-09-\d\d ", x.strip())]
        last = sept[-1].split() if sept else ["?", "", "", "", "", "", "code=?"]
        code = re.search(r"code=(\S+)", " ".join(last))
        return {"content": f"I september fejlede backuppen på NAS-2 {len(sept)} gange.\n\n"
                           f"SVAR: {last[0]}; {code.group(1) if code else '?'}; {len(sept)}"}
    if "Lav et Word-dokument, rapport.docx" in text:
        return _call("bash", command=DOCX) if done == 0 else {"content": "rapport.docx er klar."}
    if "Lav et regneark, budget.xlsx" in text:
        return _call("bash", command=XLSX) if done == 0 else {"content": "budget.xlsx er klar."}
    if "plakat til en julekoncert" in text:
        return _call("image_generate", prompt="Christmas concert poster, the headline \"JUL I HALLEN\" at the top, "
                                              "\"12. DECEMBER\" at the bottom, a lit Christmas tree, dark blue "
                                              "background") if done == 0 else {"content": "Her er plakaten."}
    if "Gør den røde bold på billedet blå" in text:
        return _call("image_edit", images=[".hub/attachments/bold.png"], prompt="Make the red ball blue, keep "
                     "everything else unchanged") if done == 0 else {"content": "Bolden er blå nu."}
    if "Fjern baggrunden fra billedet" in text:
        return _call("image_remove_background", path=".hub/attachments/bold.png") if done == 0 else \
            {"content": "Baggrunden er fjernet."}
    if "skiltet på billedet" in text:
        return {"content": "Skiltet siger ORDRE 4711. Jeg tæller fem blå cirkler og tre orange trekanter.\n\n"
                           "SVAR: ORDRE 4711; 5; 3"}
    if "Læs denne tekst højt på dansk" in text:
        quote = re.search(r'"([^"]+)"', text)
        return _call("speech_generate", text=quote.group(1) if quote else text, language="da") if done == 0 else \
            {"content": "Her er oplæsningen."}
    if "instrumentalt musikstykke på 30 sekunder" in text:
        return _call("music_generate", prompt="calm piano with strings, 80 BPM, melancholic, instrumental",
                     lyrics="", duration=30) if done == 0 else {"content": "Her er musikken."}
    if "lydeffekt på cirka 6 sekunder" in text:
        return _call("sound_generate", prompt="an old wooden door creaking slowly open, then slamming shut",
                     duration=6) if done == 0 else {"content": "Her er lyden."}
    if "rød papirdrage" in text:
        return _call("video_generate", prompt="a red paper kite flying over a green hill on a windy day, white "
                                              "clouds, the camera slowly follows the kite") if done == 0 else \
            {"content": "Her er videoen."}
    if "Få billedet til at bevæge sig" in text:
        return _call("video_generate", image=".hub/attachments/sejlbaad.png", prompt="the waves ripple and the "
                     "sailboat rocks gently, static camera") if done == 0 else {"content": "Her er videoen."}
    if "Python-modul, roman.py" in text:
        steps = [_call("write_file", path="roman.py", content=ROMAN),
                 _call("write_file", path="test_roman.py", content=ROMAN_TESTS),
                 _call("bash", command="python3 -m unittest -v test_roman")]
        return steps[done] if done < len(steps) else {"content": "roman.py og test_roman.py er skrevet; testene "
                                                                  "består."}
    if "lager/test_inventory.py fejler" in text:
        steps = [_call("write_file", path="lager/inventory.py", content=FIXED),
                 _call("bash", command="cd lager && python3 -m unittest -v")]
        return steps[done] if done < len(steps) else {
            "content": "Fire fejl: add overskrev i stedet for at lægge til, remove lod varer med 0 stå, low_stock "
                       "brugte < i stedet for <=, og parse_line beholdt mellemrum om varenummeret. Testene består."}
    if "C#-konsolprogram (.NET) i mappen ordtaeller" in text:
        steps = [_call("write_file", path="ordtaeller/ordtaeller.csproj", content=CSPROJ),
                 _call("write_file", path="ordtaeller/Program.cs", content=PROGRAM),
                 _call("bash", command="dotnet run --project ordtaeller -- tekst.txt", timeout=600)]
        if done < len(steps):
            return steps[done]
        out = [x.strip() for x in tool.splitlines() if re.match(r"^\s*[^\W\d_]+: \d+\s*$", x)]
        return {"content": "Programmet er bygget og kørt på tekst.txt:\n\n" + "\n".join(out)}
    if "interaktiv side som artefakt" in text:
        return _call("artifact", action="create", id="omregner", title="Celsius og Fahrenheit", type="html",
                     content=PAGE) if done == 0 else {"content": "Omregneren er klar."}
    if "nyeste stabile version af Python" in text:
        return _call("web_search", query="Python latest stable release") if done == 0 else \
            {"content": "Den nyeste stabile version er Python 3.15.0, udgivet 1. oktober 2026. Kilde: "
                        "https://www.python.org/downloads/"}
    if "Lav i mappen maalinger" in text:
        if done == 0:
            return _call("bash", command=FILES)
        return {"content": "Scriptet er kørt:\n\n" + "\n".join(x for x in tool.splitlines() if x.startswith("- "))}
    return None
