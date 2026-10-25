"""0.21.0 - Settings -> Tests: hard cases run as real chats and checked. The cases' own answers (the puzzle has one
solution, the fixtures hold what the checks expect), the checks on good and bad answers, a run through the agent with
the fake model (hidden chats, background priority, questions answered by the case, the sandbox's checks on what the
turn made), Stop, the time limit, the export and deleting a run."""
from __future__ import annotations

import asyncio
import io
import json
import os
import shutil
import subprocess
import sys
import zipfile
from itertools import permutations
from pathlib import Path

import pytest

from app import db
from app.evals import check_lib, fixtures
from app.evals.checks import Turn, answer_line, bullets, json_object, language, numbers, verdict
from app.evals.suite import BY_ID, CASES, LOGIC

PY = shutil.which("python3")


# ---------------------------------------------------------------- the cases themselves

def test_the_puzzle_has_one_answer():
    """The clues as LOGIC words them, solved by trying every arrangement: one solution, the one `expect` gives."""
    names, colours, pets = ("Anna", "Bent", "Cecilie", "Dennis"), ("rød", "blå", "grøn", "gul"), \
        ("hund", "kat", "papegøje", "fisk")
    clues = [lambda n, c, p: n["Anna"] == p["kat"] + 1, lambda n, c, p: c["blå"] == c["rød"] - 1,
             lambda n, c, p: p["fisk"] == c["rød"], lambda n, c, p: abs(n["Bent"] - n["Dennis"]) != 1,
             lambda n, c, p: p["papegøje"] == c["grøn"], lambda n, c, p: p["hund"] == 1,
             lambda n, c, p: abs(n["Dennis"] - p["papegøje"]) == 1]
    assert LOGIC.count("\n1. ") == 1 and "\n7. Dennis bor ved siden af" in LOGIC
    found = []
    for N in permutations(range(4)):
        n = dict(zip(names, N))
        for Co in permutations(range(4)):
            c = dict(zip(colours, Co))
            for P in permutations(range(4)):
                p = dict(zip(pets, P))
                if all(f(n, c, p) for f in clues):
                    found.append((n, c, p))
    assert len(found) == 1
    n, c, p = found[0]
    assert n["Dennis"] == p["fisk"] and c["gul"] == p["kat"]


def test_the_fixtures_hold_clear_answers():
    data, e = fixtures.sales_csv()
    assert data.decode().splitlines()[0] == "dato;region;produkt;antal;stykpris"
    assert "1.199,00" in data.decode() and e["q2_region"] == "Øst" and e["best_month"] == 11
    assert e["q2_margin"] > 1.1 and e["month_margin"] > 1.2          # no near ties
    pdf, e = fixtures.manual_pdf()
    import pymupdf
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert doc.page_count == 12
    for page, fact in fixtures.MANUAL_FACTS.items():
        assert fact in " ".join(doc[page - 1].get_text().split()), page
    log, e = fixtures.server_log()
    lines = log.decode().splitlines()
    nas2 = [x for x in lines if "NAS-2 ERROR backup FAILED" in x]
    assert [x[:10] for x in nas2 if x[5:7] == "09"] == ["2026-09-03", "2026-09-14", "2026-09-27"]
    assert e == {"date": "2026-09-27", "code": "E-5102", "count": 3, "lines": len(lines)}
    visible, hidden, e = fixtures.word_texts()
    for text, top in ((visible, e["visible"]), (hidden, e["hidden"])):
        six = fixtures.top_words(text, 6)
        assert six[:5] == top and six[4][1] > six[5][1]            # nothing ties with the fifth word
        # words that tie sort alike in any culture (C#'s default comparison is culture-aware): they differ at once,
        # in ASCII letters, and none begins with "aa" (Danish sorts that as å)
        for w, n in top:
            for x, m in top:
                if x != w and m == n:
                    assert w[0] != x[0] and w[0].isascii() and x[0].isascii() and not w.startswith("aa"), (w, x)
    for fn, size in ((fixtures.ball_png, (768, 512)), (fixtures.sign_png, (1024, 768)), (fixtures.boat_png, (1024, 576))):
        from PIL import Image
        raw, _ = fn()
        assert Image.open(io.BytesIO(raw)).size == size


def test_every_case_is_complete():
    assert len({c.id for c in CASES}) == len(CASES) >= 20
    from app.evals.suite import AREAS
    assert {c.area for c in CASES} == {a for a, _ in AREAS}
    for c in CASES:
        assert c.prompt and c.expect and c.what and c.title
        assert not c.sandbox or c.id in check_lib.CASES, c.id
        c.inputs()                                                  # makes its files


# ---------------------------------------------------------------- reading answers

def test_reading_answers():
    assert numbers("5.397,00 kr.") == [5397.0] and numbers("6 746,25") == [6.0, 746.25, 6746.25]
    assert numbers("1.800") == [1.8, 1800.0] and numbers("5,397.00") == [5397.0]
    assert answer_line("Udregning...\n**SVAR:** Dennis, gul\n", "SVAR") == "Dennis, gul"
    assert answer_line("no line", "SVAR") is None
    assert language("Kære kolleger, I er inviteret til julefrokost i kantinen.")[0] == "da"
    assert language("Dear colleagues, you are invited to the party.")[0] == "en"
    assert bullets("Hej\n- en\n* to\n1. tre\n") == ["en", "to", "tre"]
    assert json_object('```json\n{"a": 1}\n```') == ({"a": 1}, None)
    assert json_object("nope")[0] is None
    assert verdict("done", [{"ok": True, "level": "must"}, {"ok": False, "level": "should"}]) == "partial"
    assert verdict("done", [{"ok": False, "level": "must"}]) == "fail"
    assert verdict("timeout", []) == "error"


def _run_check(case_id: str, turn: Turn) -> dict:
    case = BY_ID[case_id]
    checks = case.check(turn, case.inputs().expected)
    return {c["name"]: c["ok"] for c in checks if c["level"] != "info"}


def test_the_checks_tell_good_answers_from_bad():
    good = {
        "chat.logic": "Hund i hus 2 giver ... \nSVAR: Dennis, gul",
        "chat.math": "A and B fill 5/12 in the first hour ...\nANSWER: 13:00",
        "chat.json": json.dumps({"firma": "Holm Hydraulik ApS", "kontakt": "Lene Holm", "varenummer": "KV-250",
                                 "antal": 14, "stykpris": 385.5, "total_ekskl_moms": 5397.0,
                                 "total_inkl_moms": 6746.25, "leveringsdato": "2026-11-13", "postnummer": "8600",
                                 "email": "bogholderi@holmhydraulik.dk"}),
        "chat.danish": "Kære kolleger\n\nI er inviteret til årets julefrokost:\n\n- Fredag den 4. december\n"
                       "- Vi mødes i kantinen kl. 18\n- Tilmelding senest den 20. november\n- Tag godt humør med\n\n"
                       "Venlig hilsen\nSøren",
        "documents.csv": "Beregnet med pandas.\nREGION: Øst; 408.924,50 kr.\nMÅNED: november",
        "documents.pdf": "Side 9 siger 2.150 timer, men ændringsloggen på side 11 sænker det.\nSVAR: 1.900 driftstimer",
        "documents.log": "grep viser tre fejl.\nSVAR: 2026-09-27; E-5102; 3",
        "images.vision": "SVAR: ORDRE 4711; 5; 3",
    }
    bad = {
        "chat.logic": "SVAR: Bent, rød",
        "chat.math": "ANSWER: 10:45",
        "chat.json": "Her er dataene: firma er Holm",
        "chat.danish": "Dear all, a hyggelig party on December 4.",
        "documents.csv": "REGION: Syd; 360.000 kr.\nMÅNED: juni",
        "documents.pdf": "SVAR: 2.150 driftstimer",
        "documents.log": "SVAR: 2026-10-02; E-5102; 5",
        "images.vision": "SVAR: ORDER 4117; 4; 2",
    }
    for case_id, text in good.items():
        got = _run_check(case_id, Turn(answer=text, calls=[{"name": "bash", "ok": True, "arguments": {}, "text": ""}]))
        assert all(got.values()), (case_id, got)
    for case_id, text in bad.items():
        got = _run_check(case_id, Turn(answer=text))
        assert not all(got.values()), (case_id, got)
    # the page: an artifact with two fields, a reset and the formula
    html = ("<input id=c><input id=f><button>Nulstil</button><script>f.value = c.value * 9 / 5 + 32</script>")
    ok = _run_check("code.web", Turn(artifacts=[{"id": "x", "title": "Omregner", "type": "html", "content": html}]))
    assert all(ok.values())
    assert _run_check("code.web", Turn()) == {"an artifact was made": False}
    web = _run_check("agent.web", Turn(answer="Python 3.15.0 (1. oktober 2026): https://www.python.org/downloads/",
                                       calls=[{"name": "web_search", "ok": True, "arguments": {}, "text": ""}]))
    assert all(web.values())


# ---------------------------------------------------------------- the sandbox's checks, with real files

def _check_lib(folder: Path, case_id: str, inp: dict) -> dict:
    (folder / ".hub" / "tests").mkdir(parents=True, exist_ok=True)
    (folder / ".hub" / "tests" / "input.json").write_text(json.dumps(inp))
    p = subprocess.run([PY, str(Path(check_lib.__file__)), case_id, ".hub/tests/input.json"], cwd=folder,
                       capture_output=True, text=True, timeout=600)
    line = next(x for x in p.stdout.splitlines() if x.startswith(check_lib.MARK))
    return {c["name"]: c for c in json.loads(line[len(check_lib.MARK):])}


def _has(*modules: str) -> bool:
    if not PY:
        return False
    return subprocess.run([PY, "-c", "import " + ", ".join(modules)], capture_output=True).returncode == 0


@pytest.mark.skipif(not _has("docx", "openpyxl", "PIL", "numpy"), reason="python3 without the sandbox's modules")
def test_the_sandbox_checks_read_documents_and_pictures(tmp_path):
    script = """
import docx, openpyxl
d = docx.Document()
d.add_heading("Månedsrapport oktober 2026", 0)
d.add_heading("Produktion", 1)
t = d.add_table(rows=4, cols=3)
d.add_heading("Kvalitet", 1)
d.add_paragraph("Fint.")
d.add_heading("Plan for november", 1)
for x in ("a", "b", "c"):
    d.add_paragraph(x, style="List Bullet")
d.save("rapport.docx")
wb = openpyxl.Workbook(); ws = wb.active
ws.append(["Måned", "Indtægter", "Udgifter", "Resultat"])
for i, m in enumerate(["januar","februar","marts","april","maj","juni","juli","august","september","oktober","november","december"], start=2):
    ws.append([m, 1000 + i, 500, f"=B{i}-C{i}"])
ws.append(["Total", "=SUM(B2:B13)", "=SUM(C2:C13)", "=SUM(D2:D13)"])
wb.save("budget.xlsx")
"""
    subprocess.run([PY, "-c", script], cwd=tmp_path, check=True)
    word = _check_lib(tmp_path, "documents.word", {"files": [{"path": "rapport.docx"}]})
    assert all(c["ok"] for c in word.values() if c["level"] != "info"), word
    excel = _check_lib(tmp_path, "documents.excel", {"files": [{"path": "budget.xlsx"}]})
    assert all(c["ok"] for c in excel.values() if c["level"] != "info"), excel
    # a cut-out: the ball kept, the rest transparent; an "edit" that changed nothing
    from PIL import Image, ImageDraw
    raw, geo = fixtures.ball_png()
    (tmp_path / "in.png").write_bytes(raw)
    cut = Image.new("RGBA", (768, 512), (0, 0, 0, 0))
    ImageDraw.Draw(cut).ellipse((304, 220, 464, 380), fill=(200, 30, 30, 255))
    cut.save(tmp_path / "cut.png")
    res = _check_lib(tmp_path, "images.cutout", {"images": [{"path": "cut.png"}], "expected": geo})
    assert all(c["ok"] for c in res.values() if c["level"] != "info"), res
    res = _check_lib(tmp_path, "images.edit", {"images": [{"path": "in.png"}], "inputs": {"picture": "in.png"},
                                               "expected": geo})
    assert res["the ball is blue now"]["ok"] is False and res["the rest of the picture is unchanged"]["ok"] is True


@pytest.mark.skipif(not (shutil.which("ffmpeg") and shutil.which("ffprobe") and _has("PIL", "numpy")),
                    reason="no ffmpeg")
def test_the_sandbox_checks_measure_sound_and_video(tmp_path):
    subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-f", "lavfi", "-i", "sine=f=440:d=8", "-ar", "24000",
                    str(tmp_path / "tale.wav")], check=True)
    res = _check_lib(tmp_path, "sound.speech", {"audio": [{"path": "tale.wav"}], "expected": {"seconds": [6, 30]}})
    assert res["it lasts 6-30 s"]["ok"] and res["it is not silent"]["ok"], res
    res = _check_lib(tmp_path, "sound.music", {"audio": [{"path": "tale.wav"}], "expected": {"seconds": [25, 40]}})
    assert res["it lasts 25-40 s"]["ok"] is False
    raw, _ = fixtures.boat_png()
    (tmp_path / "boat.png").write_bytes(raw)
    # the picture slowly zoomed: starts from it and moves
    subprocess.run(["ffmpeg", "-nostdin", "-loglevel", "error", "-loop", "1", "-i", "boat.png", "-vf",
                    "zoompan=z='1+0.004*on':d=1:s=1024x576:fps=24,format=yuv420p", "-t", "4", "-c:v", "libx264",
                    "video.mp4"], cwd=tmp_path, check=True)
    res = _check_lib(tmp_path, "video.i2v", {"video": [{"path": "video.mp4"}], "inputs": {"picture": "boat.png"},
                                             "expected": {"seconds": [2, 11]}})
    assert all(c["ok"] for c in res.values() if c["level"] != "info"), res


# ---------------------------------------------------------------- runs through the agent

FIXED = fixtures.INVENTORY.replace("self.items[sku] = qty\n", "self.items[sku] = self.items.get(sku, 0) + qty\n") \
    .replace("self.items[sku] = have - qty\n", "self.items[sku] = have - qty\n        if not self.items[sku]:\n"
                                               "            del self.items[sku]\n") \
    .replace("if qty < threshold", "if qty <= threshold").replace("return sku, int(qty)", "return sku.strip(), int(qty)")


def _user_text(body: dict) -> str:
    for m in body.get("messages") or []:
        if m.get("role") == "user":
            c = m.get("content")
            return c if isinstance(c, str) else " ".join(p.get("text", "") for p in c if p.get("type") == "text")
    return ""


def _tool_results(body: dict) -> list[dict]:
    return [m for m in body.get("messages") or [] if m.get("role") == "tool"]


def responder(fix: bool = True, slow: float = 0.0):
    """The fake model: answers the cases it knows - the puzzle right, the tank wrong, the stock register by writing
    the fixed file (or not), the web case by searching first."""
    def reply(body: dict) -> dict:
        text, done = _user_text(body), _tool_results(body)
        if "Fire personer" in text:
            if not done:
                return {"tool_calls": [{"name": "web_search", "arguments": {"query": "zebra puzzle"}}]}
            return {"content": "Hunden bor i hus 2 ...\nSVAR: Dennis, gul", "delay": slow}
        if "A tank is filled" in text:
            return {"content": "ANSWER: 10:45", "delay": slow}
        if "lager/test_inventory.py" in text:
            if not done:
                return {"tool_calls": [{"name": "write_file", "arguments": {
                    "path": "lager/inventory.py", "content": FIXED if fix else fixtures.INVENTORY + "\n# looked\n"}}]}
            if len(done) == 1:
                return {"tool_calls": [{"name": "bash", "arguments": {"command": "cd lager && python3 -m unittest"}}]}
            return {"content": "Fire fejl rettet." if fix else "Jeg kiggede på det."}
        if "nyeste stabile version af Python" in text:
            if not done:
                return {"tool_calls": [{"name": "web_search", "arguments": {"query": "python latest release"}}]}
            return {"content": "Python 3.15.0 udkom 1. oktober 2026: https://www.python.org/downloads/"}
        return {"content": "Hello from the fake model."}
    return reply


def local_sandbox(workspace: Path):
    """The fake sandbox runs the checks' command for real, in the case's folder (python3 with the sandbox's modules)."""
    def handler(body: dict) -> dict:
        cwd = workspace / str(body.get("cwd") or "/workspace").removeprefix("/workspace").lstrip("/")
        if not body["cmd"].startswith("python3 .hub/tests/check_lib.py"):
            return {"exit_code": 0, "stdout": "", "stderr": "", "timed_out": False, "duration_ms": 1}
        p = subprocess.run(["bash", "-c", body["cmd"].replace("python3", PY or sys.executable, 1)], cwd=cwd,
                           capture_output=True, text=True, timeout=300, env={**os.environ, **(body.get("env") or {})})
        return {"exit_code": p.returncode, "stdout": p.stdout, "stderr": p.stderr, "timed_out": False,
                "duration_ms": 5, "truncated": False}
    return handler


async def _finish(client, run_id: str, timeout: float = 30) -> dict:
    for _ in range(int(timeout * 20)):
        r = await client.get(f"/hub/api/tests/runs/{run_id}")
        if r.json()["status"] != "running":
            return r.json()
        await asyncio.sleep(0.05)
    raise AssertionError("the test run did not finish")


async def test_the_overview_says_what_each_case_needs(signed_in):
    r = await signed_in.get("/hub/api/tests")
    assert r.status_code == 200, r.text
    o = r.json()
    cases = {c["id"]: c for c in o["cases"]}
    assert len(cases) == len(CASES) and [a["id"] for a in o["areas"]][0] == "chat"
    assert cases["images.poster"]["ready"] is True                         # Qwen-Image 2.1 (conftest)
    assert cases["video.t2v"]["ready"] is False and cases["video.t2v"]["missing"] == ["video.t2v"]
    assert cases["chat.logic"]["ready"] is True and o["running"] is None and o["runs"] == []
    assert {"qwen3.8-27b", "qwen3.6-35b-a3b"} <= {m["id"] for m in o["models"]}


async def test_a_run_checks_answers_hides_its_chats_and_lets_chats_go_first(signed_in, app, fakes):
    fakes.llama["responder"] = responder()
    r = await signed_in.post("/hub/api/tests/runs", json={"cases": ["chat.logic", "chat.math", "video.t2v"],
                                                          "model": "qwen3.6-35b-a3b", "thinking": False})
    assert r.status_code == 200, r.text
    run_id = r.json()["id"]
    again = await signed_in.post("/hub/api/tests/runs", json={"cases": ["chat.logic"]})
    assert again.status_code == 409
    rec = await _finish(signed_in, run_id)
    assert rec["status"] == "done" and rec["model"] == "qwen3.6-35b-a3b"
    res = {x["id"]: x for x in rec["results"]}
    assert res["chat.logic"]["status"] == "pass", res["chat.logic"]["checks"]
    assert res["chat.logic"]["approvals"] == [{"tool": "web_search", "what": "search", "asked": "zebra puzzle",
                                               "answer": "deny", "why": "a test answers No to anything it does not need"}] \
        or res["chat.logic"]["approvals"] == []          # web search asks only when the internet asks first
    assert res["chat.math"]["status"] == "fail"
    assert res["video.t2v"]["status"] == "skipped" and "video.t2v" in res["video.t2v"]["why"]
    assert res["chat.logic"]["model"] == "qwen3.6-35b-a3b" and res["chat.logic"]["seconds"] > 0
    # the chats are there, but not in the chat list
    listed = {c["id"] for c in (await signed_in.get("/hub/api/chats")).json()}
    assert res["chat.logic"]["chat_id"] not in listed
    chat = (await signed_in.get(f"/hub/api/chats/{res['chat.logic']['chat_id']}")).json()["chat"]
    assert chat["title"] == "Test: Logic puzzle" and chat["test_run"] == run_id
    # its turns queued as a test run's, behind the chats
    with db.connect() as con:
        jobs = con.execute("SELECT owner, priority FROM jobs_log WHERE owner='tests'").fetchall()
    assert jobs and all(j["priority"] == "background" for j in jobs)
    # thinking off as asked
    sent = [b for b in fakes.llama["requests"] if "Fire personer" in _user_text(b)]
    assert sent and all(b.get("chat_template_kwargs", {}).get("enable_thinking") is False for b in sent
                        if "chat_template_kwargs" in b)
    o = (await signed_in.get("/hub/api/tests")).json()
    assert o["runs"][0]["counts"]["pass"] == 1 and o["runs"][0]["counts"]["skipped"] == 1
    assert {c["id"]: c["last_seconds"] for c in o["cases"]}["chat.logic"] > 0


async def test_questions_are_answered_as_the_case_allows(signed_in, fakes):
    await signed_in.put("/hub/api/settings", json={"section": "web", "values": {"access": "ask"}})
    fakes.llama["responder"] = responder()
    r = await signed_in.post("/hub/api/tests/runs", json={"cases": ["agent.web", "chat.logic"]})
    rec = await _finish(signed_in, r.json()["id"])
    res = {x["id"]: x for x in rec["results"]}
    web, logic = res["agent.web"], res["chat.logic"]
    assert [a["answer"] for a in web["approvals"]] == ["allow"] and web["status"] == "pass", web["checks"]
    assert web["tools"] == [{"name": "web_search", "ok": True}]
    assert [a["answer"] for a in logic["approvals"]] == ["deny"]
    assert logic["tools"] == [{"name": "web_search", "ok": False}] and logic["status"] == "pass"


@pytest.mark.skipif(not PY, reason="no python3")
async def test_the_sandbox_checks_what_the_turn_made(signed_in, app, fakes, hub_env):
    fakes.sandbox["handler"] = local_sandbox(hub_env.workspace_dir)
    fakes.llama["responder"] = responder(fix=True)
    rec = await _finish(signed_in, (await signed_in.post("/hub/api/tests/runs", json={"cases": ["code.debug"]})).json()["id"])
    good = rec["results"][0]
    checks = {c["name"]: c for c in good["checks"]}
    assert good["status"] == "pass", good["checks"]
    assert checks["the given tests pass"]["detail"] == "5 of 5 passed"
    assert checks["hidden tests pass"]["ok"] and checks["the tests were not changed"]["ok"]
    assert good["files"] == ["lager/inventory.py"]
    fakes.llama["responder"] = responder(fix=False)
    rec = await _finish(signed_in, (await signed_in.post("/hub/api/tests/runs", json={"cases": ["code.debug"]})).json()["id"])
    bad = {c["name"]: c for c in rec["results"][0]["checks"]}
    assert rec["results"][0]["status"] == "fail"
    assert bad["the given tests pass"]["ok"] is False and "passed" in bad["the given tests pass"]["detail"]
    # the sandbox ran the checks offline where it can, in the case's folder
    run = [e for e in fakes.sandbox["execs"] if "check_lib.py" in e["cmd"]]
    assert run and run[-1]["cwd"].startswith("/workspace/chats/") and run[-1].get("offline") is True


async def test_stop_and_the_time_limit(signed_in, app, fakes, monkeypatch):
    fakes.llama["responder"] = responder(slow=5)
    r = await signed_in.post("/hub/api/tests/runs", json={"cases": ["chat.logic", "chat.math"]})
    run_id = r.json()["id"]
    for _ in range(100):
        live = (await signed_in.get(f"/hub/api/tests/runs/{run_id}")).json()["live"]
        if live and live["chat_id"]:
            break
        await asyncio.sleep(0.05)
    assert live["case"] == "chat.logic"
    assert (await signed_in.post(f"/hub/api/tests/runs/{run_id}/stop")).status_code == 200
    rec = await _finish(signed_in, run_id)
    assert rec["status"] == "stopped" and [x["status"] for x in rec["results"]] == ["stopped"]
    assert (await signed_in.post(f"/hub/api/tests/runs/{run_id}/stop")).status_code == 409
    monkeypatch.setattr(BY_ID["chat.math"], "timeout", 1)
    rec = await _finish(signed_in, (await signed_in.post("/hub/api/tests/runs", json={"cases": ["chat.math"]})).json()["id"])
    assert rec["results"][0]["status"] == "error" and rec["results"][0]["why"] == "The case's time ran out"


async def test_the_export_and_deleting_a_run(signed_in, app, fakes, hub_env):
    fakes.llama["responder"] = responder()
    rec = await _finish(signed_in, (await signed_in.post("/hub/api/tests/runs",
                                                         json={"cases": ["chat.logic", "video.t2v"]})).json()["id"])
    r = await signed_in.get(f"/hub/api/tests/runs/{rec['id']}/export")
    assert r.status_code == 200 and r.headers["content-type"] == "application/zip"
    z = zipfile.ZipFile(io.BytesIO(r.content))
    names = z.namelist()
    base = names[0].split("/")[0] + "/"
    for want in ("README.md", "summary.md", "results.json", "cases/01-chat.logic/case.json",
                 "cases/01-chat.logic/answer.md", "cases/01-chat.logic/events.json",
                 "cases/01-chat.logic/chat/chat.json", "cases/01-chat.logic/chat/transcript.md",
                 "cases/02-video.t2v/case.json"):
        assert base + want in names, (want, names)
    summary = z.read(base + "summary.md").decode()
    assert "| 1 | Chat | Logic puzzle | pass |" in summary and "SVAR: Dennis, gul" in summary
    assert "A good answer: Dennis has the fish" in summary and "skipped" in summary
    assert not list((hub_env.data_dir / "tmp").glob("tests-*.zip"))           # the file went once it was sent
    chat_id = rec["results"][0]["chat_id"]
    folder = hub_env.workspace_dir / "chats" / str(chat_id)
    assert folder.is_dir()
    assert (await signed_in.delete(f"/hub/api/tests/runs/{rec['id']}")).status_code == 200
    assert (await signed_in.get(f"/hub/api/chats/{chat_id}")).status_code == 404 and not folder.exists()
    assert (await signed_in.get(f"/hub/api/tests/runs/{rec['id']}")).status_code == 404


async def test_a_run_cut_short_by_a_restart_says_so(app):
    with db.connect() as con:
        con.execute("INSERT INTO test_runs (id, created_at, status) VALUES ('test_x', '2026-10-25T10:00:00+00:00', "
                    "'running')")
    app.state.hub.extras["tests"].interrupted()
    assert app.state.hub.extras["tests"].get("test_x")["status"] == "interrupted"


async def test_a_person_cannot_run_tests(signed_in, app):
    from .test_people import add_person, client_from, sign_in
    await add_person(signed_in)
    async with client_from(app) as phone:
        await sign_in(phone)
        assert (await phone.get("/hub/api/tests")).status_code == 403
        assert (await phone.post("/hub/api/tests/runs", json={"cases": ["chat.logic"]})).status_code == 403
